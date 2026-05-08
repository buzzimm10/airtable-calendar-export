import logging
from datetime import datetime, date, timedelta
from typing import Dict, List, Optional, Tuple

from copy import copy
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.cell.cell import Cell, MergedCell

from config.layouts import TARGET_SHEET
from core.dates import parse_date

def get_writable_cell(ws, row: int, col: int, unmerge_if_needed: bool = False):
    """
    Return a writable Cell for a target position.
    If the position is inside a merged range, return the top-left cell.
    If the top-left is somehow still a MergedCell (can happen after insert_cols),
    optionally unmerge that range and return a normal writable cell.
    """
    for cr in list(ws.merged_cells.ranges):
        if cr.min_row <= row <= cr.max_row and cr.min_col <= col <= cr.max_col:
            top_left = ws.cell(row=cr.min_row, column=cr.min_col)

            if isinstance(top_left, MergedCell):
                if unmerge_if_needed:
                    ws.unmerge_cells(str(cr))
                    return ws.cell(row=cr.min_row, column=cr.min_col)
                raise RuntimeError(
                    f"Top-left of merged range {cr} is still a MergedCell at "
                    f"({cr.min_row}, {cr.min_col})"
                )

            return top_left

    cell = ws.cell(row=row, column=col)

    if isinstance(cell, MergedCell):
        if unmerge_if_needed:
            for cr in list(ws.merged_cells.ranges):
                if cr.min_row <= row <= cr.max_row and cr.min_col <= col <= cr.max_col:
                    ws.unmerge_cells(str(cr))
                    return ws.cell(row=cr.min_row, column=cr.min_col)
        raise RuntimeError(f"Cell {get_column_letter(col)}{row} is a MergedCell and could not be resolved.")

    return cell

def ensure_writable_cell(ws, row: int, col: int):
    """Return a real writable cell after unmerging any merged range covering the target."""
    merged_ranges = [cr for cr in list(ws.merged_cells.ranges)
                     if cr.min_row <= row <= cr.max_row and cr.min_col <= col <= cr.max_col]
    logging.debug(
        "ensure_writable_cell: target %s%d merged_ranges=%s",
        get_column_letter(col), row, merged_ranges,
    )
    for cr in merged_ranges:
        logging.debug(
            "ensure_writable_cell: target %s%d in merged range %s, unmerging",
            get_column_letter(col), row, cr,
        )
        ws.unmerge_cells(str(cr))

    cell = ws.cell(row=row, column=col)
    logging.debug(
        "ensure_writable_cell: after unmerge target %s%d cell_type=%s",
        get_column_letter(col), row, type(cell),
    )
    if isinstance(cell, MergedCell):
        logging.debug(
            "ensure_writable_cell: still MergedCell at %s%d after unmerge, returning top-left of active merged range",
            get_column_letter(col), row,
        )
        for cr in list(ws.merged_cells.ranges):
            if cr.min_row <= row <= cr.max_row and cr.min_col <= col <= cr.max_col:
                top_left = ws.cell(row=cr.min_row, column=cr.min_col)
                if isinstance(top_left, MergedCell):
                    logging.debug(
                        "ensure_writable_cell: top-left still MergedCell for range %s, unmerging it",
                        cr,
                    )
                    ws.unmerge_cells(str(cr))
                    return ws.cell(row=cr.min_row, column=cr.min_col)
                logging.debug(
                    "ensure_writable_cell: using top-left %s%d for merged range %s",
                    get_column_letter(cr.min_col), cr.min_row, cr,
                )
                return top_left

        logging.debug(
            "ensure_writable_cell: target %s%d is a stale MergedCell with no active merged range, replacing with fresh Cell",
            get_column_letter(col), row,
        )
        real_cell = Cell(ws, row=row, column=col)
        ws._cells[(row, col)] = real_cell
        return real_cell

    return cell

def cleanup_stale_merged_cells(ws, row_start: int = 1, row_end: Optional[int] = None, col_start: int = 1, col_end: Optional[int] = None):
    """
    Replace stale MergedCell objects that are no longer covered by an active merged range
    with normal writable Cell objects.
    """
    if row_end is None:
        row_end = ws.max_row
    if col_end is None:
        col_end = ws.max_column

    active_ranges = list(ws.merged_cells.ranges)

    def is_inside_active_merge(r: int, c: int) -> bool:
        for cr in active_ranges:
            if cr.min_row <= r <= cr.max_row and cr.min_col <= c <= cr.max_col:
                return True
        return False

    replaced = 0

    for row in range(row_start, row_end + 1):
        for col in range(col_start, col_end + 1):
            key = (row, col)
            cell_obj = ws._cells.get(key)

            if isinstance(cell_obj, MergedCell) and not is_inside_active_merge(row, col):
                ws._cells[key] = Cell(ws, row=row, column=col)
                replaced += 1

    if replaced:
        logging.info(
            "cleanup_stale_merged_cells: replaced %s stale merged cells in range %s:%s to %s:%s",
            replaced,
            get_column_letter(col_start), row_start,
            get_column_letter(col_end), row_end,
        )


def safe_set_value(ws, row: int, col: int, value, alignment=None, number_format=None, font=None, fill=None):
    cell = ensure_writable_cell(ws, row, col)
    cell.value = value
    if alignment is not None:
        cell.alignment = alignment
    if number_format is not None:
        cell.number_format = number_format
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    return cell

def force_unmerged_writable_cell(ws, row: int, col: int):
    """
    Unmerge any merged range that contains the target cell, then return a normal writable cell.
    This is intentionally aggressive and is meant for dynamic header/deadline cells after insert_cols().
    """
    return ensure_writable_cell(ws, row, col)

def cell_value(ws, row: int, col: int):
    return ws.cell(row=row, column=col).value

def center_wrap():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)


def left_wrap():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

def copy_cell_style(src, dst):
    if src.has_style:
        dst._style = copy(src._style)
    dst.number_format = src.number_format
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.protection = copy(src.protection)

def get_anchor_dates_from_values_wb(template_path: str, anchors: List[str]) -> Dict[str, date]:
    wb_vals = load_workbook(template_path, data_only=True)
    if TARGET_SHEET not in wb_vals.sheetnames:
        raise RuntimeError(f"Target sheet '{TARGET_SHEET}' not found for reading anchors.")
    ws_vals = wb_vals[TARGET_SHEET]
    out = {}
    for addr in anchors:
        val = ws_vals[addr].value
        d = parse_date(val)
        if not d:
            logging.warning(
                f"Anchor {addr} value unreadable (got {val}). Open template in Excel to recalc/save, or pass --anchor-dates."
            )
        else:
            out[addr] = d
    return out

def _to_indices(layout_letters):
    d = {}
    for day, cols in layout_letters.items():
        d[day] = [(role, column_index_from_string(letter)) for role, letter in cols]
    return d
