import logging
from copy import copy
from typing import Dict, List, Optional
from datetime import date, timedelta
from collections import defaultdict, Counter

from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

from config.fields import AT_FIELDS
from config.layouts import DAY_ORDER, DAY_LAYOUT
from config.styles import OUT_OF_MONTH_FILL
from config.countries import SHEET_COUNTRIES

from core.text_formatting import normalize_str, lower_clean, replace_month_in_title
from core.dates import parse_date, month_name_from_records
from core.excel_utils import safe_set_value, get_writable_cell, cell_value
from core.countries import collect_included_countries

from openpyxl.styles import Alignment
from core.text_formatting import process_multiline_branding, markdown_to_richtext
from config.fields import AT_FIELDS

from services.planner_service import (
    WeekBlock,
    hub_role,
    asset_is_event,
    earliest_deadline_for_day,
    assemble_text_for_single_asset,
)


def rebuild_full_calendar_week_merges(ws, block: WeekBlock):
    """
    Rebuild merges for:
    - WEEK # row: B:C
    - Live Date row: merged across each day group
    - Deadline row: merged across each day group
    """

    # 1) limpiar merges horizontales existentes dentro de las filas relevantes
    rows_to_clean = {block.anchor_row}
    if block.deadline_row:
        rows_to_clean.add(block.deadline_row)

    for cr in list(ws.merged_cells.ranges):
        if (
            cr.min_row == cr.max_row
            and (
                cr.min_row in rows_to_clean
                or (cr.min_row == block.anchor_row)
            )
        ):
            ws.unmerge_cells(str(cr))

    # 2) WEEK # -> B:C
    ws.merge_cells(
        start_row=block.anchor_row,
        start_column=2,
        end_row=block.anchor_row,
        end_column=3,
    )

    # 3) Live Date y Deadline por cada día
    for day_name in DAY_ORDER:
        day_cols = [col for _, col in DAY_LAYOUT[day_name]]
        start_col = min(day_cols)
        end_col = max(day_cols)

        # Live Date
        if end_col > start_col:
            ws.merge_cells(
                start_row=block.anchor_row,
                start_column=start_col,
                end_row=block.anchor_row,
                end_column=end_col,
            )

        # Deadline
        if block.deadline_row and end_col > start_col:
            ws.merge_cells(
                start_row=block.deadline_row,
                start_column=start_col,
                end_row=block.deadline_row,
                end_column=end_col,
            )

def write_to_day_role_columns(ws, block, day_name, role, records, enable_mediaplanning):
    if role == "Mediaplanning" and not enable_mediaplanning:
        return

    role_cols = [col for r, col in DAY_LAYOUT[day_name] if r == role]
    uk_cols = [col for r, col in DAY_LAYOUT[day_name] if r == "UK"]

    def _write(row_idx, col_idx, value):
        if not row_idx:
            return
        safe_set_value(
            ws,
            row_idx,
            col_idx,
            value,
            alignment=Alignment(horizontal='center', vertical='center', wrap_text=True),
        )

    branding_layout_dict = None
    if block.tc_row:
        branding_layout_dict = {
            "watermark": block.tc_row + 1,
            "bug": [block.tc_row + 2, block.tc_row + 3],
            "endcard": [block.tc_row + 4, block.tc_row + 5]
        }
        
    if role == "Programming":
        for idx, record in enumerate(records[:2]):
            col = role_cols[idx]
            single = assemble_text_for_single_asset(record)

            _write(block.content_type_row, col, single["content_type"])
            _write(block.property_row, col, single["property"])
            _write(block.title_row, col, single["title"])
            _write(block.synopsis_row, col, single["synopsis"])
            _write(block.strategy_row, col, single["strategy"])
            _write(block.tx_row, col, single["tx"])
            _write(block.tc_row, col, single["tc"])
            
            # --- INYECTOR DE NOTAS DE DERECHOS (NEW) ---
            fields = record.get("fields", record)
            policy = fields.get("Rights Management Policy")
            policy_str = str(policy).strip() if policy else ""
            extra_note = ""
            
            if "Block AUS/NZ" in policy_str:
                extra_note = "\n\n⚠️ **Rights Policy:** Block AUS/NZ/TAIWAN/CHINA/HONG KONG - Monetize everywhere else"
            elif "Monetize Locally" in policy_str:
                extra_note = "\n\n⚠️ **Rights Policy:** Monetize Locally (Geo-block)"
            elif "Block Everywhere" in policy_str:
                extra_note = "\n\n🚫 **Rights Policy:** Block Everywhere / Do Not Monetize"
                
            comments_text = single["comments"] + extra_note
            _write(block.comments_row, col, markdown_to_richtext(comments_text))

            if branding_layout_dict:
                write_dc_branding(ws, branding_layout_dict, col, record)

    elif role == "Mediaplanning":
        target_cols = []
        if role_cols: target_cols.append(role_cols[0])
        if uk_cols: target_cols.append(uk_cols[0])

        num_records = len(records[:2])
        
        if num_records > 1 and len(target_cols) == 2 and block.owner_row:
            safe_set_value(ws, block.owner_row, target_cols[1], "MEDIAPLANNING HUB", alignment=Alignment(horizontal='center', vertical='center', wrap_text=True))

        for idx, record in enumerate(records[:2]):
            if idx >= len(target_cols): break
            col = target_cols[idx]
            single = assemble_text_for_single_asset(record)

            _write(block.content_type_row, col, single["content_type"])
            _write(block.property_row, col, single["property"])
            
            text_to_merge = single["title"]
            if single["synopsis"]: text_to_merge += f"\n\n{single['synopsis']}"
            if single["strategy"]: text_to_merge += f"\n\n{single['strategy']}"
            
            _write(block.title_row, col, text_to_merge)
            safe_set_value(ws, block.synopsis_row, col, None)
            safe_set_value(ws, block.strategy_row, col, None)

            _write(block.tx_row, col, single["tx"])
            _write(block.tc_row, col, single["tc"])
            
            # --- NUEVO: BURNED-IN TEXT (SPAIN ONLY) ---
            if hasattr(block, 'burned_in_text_row') and block.burned_in_text_row:
                burned_in = ""
                # Solo aplica si es Vertical Short (ignorando mayúsculas/minúsculas)
                if single["content_type"].lower() == "vertical short":
                    fields = record.get("fields", record)
                    # Usamos la llave directa de Airtable
                    burned_in = str(fields.get("Burned-in Text (Spain Only)", "")).strip()
                _write(block.burned_in_text_row, col, burned_in)
            
            # --- INYECTOR DE NOTAS DE DERECHOS (NEW) ---
            fields = record.get("fields", record)
            policy = fields.get("Rights Management Policy")
            policy_str = str(policy).strip() if policy else ""
            extra_note = ""
            
            if "Block AUS/NZ" in policy_str:
                extra_note = "\n\n⚠️ **Rights Policy:** Block AUS/NZ/TAIWAN/CHINA/HONG KONG - Monetize everywhere else"
            elif "Monetize Locally" in policy_str:
                extra_note = "\n\n⚠️ **Rights Policy:** Monetize Locally (Geo-block)"
            elif "Block Everywhere" in policy_str:
                extra_note = "\n\n🚫 **Rights Policy:** Block Everywhere / Do Not Monetize"
                
            comments_text = single["comments"] + extra_note
            _write(block.comments_row, col, markdown_to_richtext(comments_text))

            if branding_layout_dict: write_dc_branding(ws, branding_layout_dict, col, record)

            if block.title_row and block.strategy_row:
                try:
                    ws.merge_cells(start_row=block.title_row, start_column=col, end_row=block.strategy_row, end_column=col)
                except Exception: pass
                cell = ws.cell(row=block.title_row, column=col)
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                
def write_week_label(ws, block, week_number_value: str):
    """
    Write 'WEEK <N>' into B row of the anchor line.
    Safe version for dynamic layout.
    """
    safe_set_value(
        ws,
        block.anchor_row,
        2,
        f"WEEK {week_number_value}",
        alignment=Alignment(horizontal='center', vertical='center', wrap_text=True),
    )

def rewrite_full_calendar_owner_row(ws, block: WeekBlock, day_layout):
    """
    Fill OWNER row according to fixed columns in each day.
    """
    if not block.owner_row:
        return

    for day_name in DAY_ORDER:
        prog_cols = [col for role, col in day_layout[day_name] if role == "Programming"]
        media_cols = [col for role, col in day_layout[day_name] if role == "Mediaplanning"]
        uk_cols = [col for role, col in day_layout[day_name] if role == "UK"]

        for col in prog_cols:
            safe_set_value(
                ws,
                block.owner_row,
                col,
                "PROGRAMMING HUB",
                alignment=Alignment(horizontal='center', vertical='center', wrap_text=True),
            )

        for col in media_cols:
            safe_set_value(
                ws,
                block.owner_row,
                col,
                "MEDIAPLANNING HUB",
                alignment=Alignment(horizontal='center', vertical='center', wrap_text=True),
            )

        for col in uk_cols:
            safe_set_value(
                ws,
                block.owner_row,
                col,
                "YT TEAM UK",
                alignment=Alignment(horizontal='center', vertical='center', wrap_text=True),
            )

def shade_out_of_month_days_full_calendar(ws, blocks: Dict[date, WeekBlock], target_year: int, target_month: int):
    grey_fill = PatternFill(fill_type="solid", fgColor=OUT_OF_MONTH_FILL)

    for block in blocks.values():
        for idx, day_name in enumerate(DAY_ORDER):
            this_date = block.monday_date + timedelta(days=idx)

            if this_date.year == target_year and this_date.month == target_month:
                continue

            day_cols = [col for _, col in DAY_LAYOUT[day_name]]

            # no pintar las 2 últimas filas del bloque
            shade_until_row = max(block.anchor_row, block.end_row - 2)

            for col in day_cols:
                for row in range(block.anchor_row, shade_until_row + 1):
                    cell = ws.cell(row=row, column=col)
                    cell.fill = copy(grey_fill)

def set_full_calendar_column_widths(ws):
    for day_name in DAY_ORDER:
        prog_cols = [col for role, col in DAY_LAYOUT[day_name] if role == "Programming"]
        media_cols = [col for role, col in DAY_LAYOUT[day_name] if role == "Mediaplanning"]
        uk_cols = [col for role, col in DAY_LAYOUT[day_name] if role == "UK"]

        for col in prog_cols:
            ws.column_dimensions[get_column_letter(col)].width = 40

        for col in media_cols:
            ws.column_dimensions[get_column_letter(col)].width = 18

        for col in uk_cols:
            ws.column_dimensions[get_column_letter(col)].width = 14

def auto_adjust_full_calendar_row_heights(ws, blocks):
    for block in blocks.values():
        row_heights = {
            block.property_row: 28,
            block.content_type_row: 24,
            block.title_row: 42,
            block.synopsis_row: 72,
            block.strategy_row: 36,
            block.tx_row: 30,
            block.tc_row: 36,
            block.burned_in_text_row: 30,
            block.comments_row: 48,
        }

        for row_idx, height in row_heights.items():
            if row_idx:
                ws.row_dimensions[row_idx].height = max(
                    ws.row_dimensions[row_idx].height or 15,
                    height
                )

def write_week_deadlines_fixed(ws, block, per_day_records):
    if not block.deadline_row:
        return

    for day_name in DAY_ORDER:
        recs = per_day_records.get(day_name, [])
        if not recs:
            continue

        d = earliest_deadline_for_day(recs)
        if d:
            first_col = [col for _, col in DAY_LAYOUT[day_name]][0]
            safe_set_value(
                ws,
                block.deadline_row,
                first_col,
                d,
                alignment=Alignment(horizontal='center', vertical='center', wrap_text=True),
                number_format="dddd, dd mmmm yyyy",
            )

def write_countries_rollout_fixed(ws, block, per_day_role_records, channel_id_to_name_map=None):
    if channel_id_to_name_map is None:
        channel_id_to_name_map = {}

    for day_name in DAY_ORDER:
        # Programming: col1 y col2
        prog_cols = [col for r, col in DAY_LAYOUT[day_name] if r == "Programming"]
        prog_recs = per_day_role_records.get((day_name, "Programming"), [])

        for idx, record in enumerate(prog_recs[:2]):
            target_col = prog_cols[idx]
            included, _ = collect_included_countries(
                youtube_channel=record.get(AT_FIELDS["youtube_channel"]),
                region=record.get(AT_FIELDS["region"]),
                notes=record.get(AT_FIELDS["notes"]),
                excluded_markets=record.get(AT_FIELDS["excluded_markets"]),
                channel_id_to_name_map=channel_id_to_name_map,
            )

            for country_name, rollout_row in block.rollout_country_rows.items():
                val = "YES" if country_name in included else "NO"
                safe_set_value(
                    ws, rollout_row, target_col, val,
                    alignment=Alignment(horizontal='center', vertical='center', wrap_text=True),
                    font=Font(color="FFFF0000" if (val == "NO" or "italy" in country_name.lower()) else "FF000000"),
                )

            for country_name, status_row in block.status_country_rows.items():
                val = "YES" if country_name in included else "NO"
                safe_set_value(
                    ws, status_row, target_col, val,
                    alignment=Alignment(horizontal='center', vertical='center', wrap_text=True),
                    font=Font(color="FFFF0000" if (val == "NO" or "italy" in country_name.lower()) else "FF000000"),
                )

        # Mediaplanning: col3 y opcional col4
        media_recs = per_day_role_records.get((day_name, "Mediaplanning"), [])
        media_cols = [col for r, col in DAY_LAYOUT[day_name] if r == "Mediaplanning"]
        uk_cols = [col for r, col in DAY_LAYOUT[day_name] if r == "UK"]
        target_cols = []
        if media_cols:
            target_cols.append(media_cols[0])
        if uk_cols:
            target_cols.append(uk_cols[0])

        for idx, record in enumerate(media_recs[:2]):
            if idx >= len(target_cols):
                break

            target_col = target_cols[idx]
            included, _ = collect_included_countries(
                youtube_channel=record.get(AT_FIELDS["youtube_channel"]),
                region=record.get(AT_FIELDS["region"]),
                notes=record.get(AT_FIELDS["notes"]),
                excluded_markets=record.get(AT_FIELDS["excluded_markets"]),
                channel_id_to_name_map=channel_id_to_name_map,
            )

            # --- INICIO DEL INTERCEPTOR: REGLA DE PORTUGAL ---
            ct_val = str(record.get("fields", record).get(AT_FIELDS["content_type"], "")).upper()
            region_val = str(record.get("fields", record).get(AT_FIELDS["region"], "")).upper()
            
            if "VERTICAL SHORT" in ct_val and "TIER A" in region_val:
                included.add("Portugal")
            # --- FIN DEL INTERCEPTOR ---

            for country_name, rollout_row in block.rollout_country_rows.items():
                val = "YES" if country_name in included else "NO"
                safe_set_value(
                    ws, rollout_row, target_col, val,
                    alignment=Alignment(horizontal='center', vertical='center', wrap_text=True),
                    font=Font(color="FFFF0000" if (val == "NO" or "italy" in country_name.lower()) else "FF000000"),
                )

            for country_name, status_row in block.status_country_rows.items():
                val = "YES" if country_name in included else "NO"
                safe_set_value(
                    ws, status_row, target_col, val,
                    alignment=Alignment(horizontal='center', vertical='center', wrap_text=True),
                    font=Font(color="FFFF0000" if (val == "NO" or "italy" in country_name.lower()) else "FF000000"),
                )

def write_full_calendar_start_date(ws, start_monday: date):
    """
    Only writes D6. The rest of the sheet updates via Excel formulas.
    """
    safe_set_value(
        ws,
        6,
        column_index_from_string("D"),
        start_monday,
        alignment=Alignment(horizontal='center', vertical='center', wrap_text=True),
        number_format="dddd, dd mmmm yyyy",
    )

from core.text_formatting import process_multiline_branding
from core.excel_utils import safe_set_value, center_wrap
from config.fields import AT_FIELDS

def write_dc_branding(ws, branding_layout, col, record):
    """
    Inyecta la sección BRANDING en DC basada en posicionamiento relativo.
    """
    f = record.get("fields", record)
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    def _set_rich_cell(row, column, text_val, alignment):
        cell = ws.cell(row=row, column=column)
        cell.value = markdown_to_richtext(text_val)
        cell.alignment = alignment

    # 1. WATERMARK: 1 fila con todo el bloque de texto (\n interno y Markdown)
    watermark_text = f.get(AT_FIELDS["branding_watermark"], "")
    _set_rich_cell(branding_layout["watermark"], col, watermark_text, align_center)

    # 2. BUG: 2 filas independientes
    bug_lines = process_multiline_branding(f.get(AT_FIELDS["branding_bug"]), max_lines=2)
    for idx, row in enumerate(branding_layout["bug"]):
        _set_rich_cell(row, col, bug_lines[idx], align_center)

    # 3. ENDCARD: 2 filas independientes
    endcard_lines = process_multiline_branding(f.get(AT_FIELDS["branding_endcard"]), max_lines=2)
    for idx, row in enumerate(branding_layout["endcard"]):
        _set_rich_cell(row, col, endcard_lines[idx], align_center)

def update_full_calendar_month_title(ws, records: List[Dict]):
    """
    Update B3:Z3 merged title in FULL CALENDAR replacing MONTH with actual month.
    """
    month_name_upper = month_name_from_records(records)
    cell = get_writable_cell(ws, 3, 2)  # B3
    existing = normalize_str(cell.value)
    if existing:
        cell.value = replace_month_in_title(existing, month_name_upper)
    else:
        cell.value = month_name_upper
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def find_week_blocks(ws_write, anchor_dates: Dict[str, date]) -> Dict[date, WeekBlock]:
    anchors = []
    for addr, mon_date in anchor_dates.items():
        cell = ws_write[addr]
        anchors.append((cell.row, mon_date, addr))
    anchors = [a for a in anchors if a[1] is not None]
    if not anchors:
        raise RuntimeError("No valid week anchors resolved.")
    anchors_sorted = sorted(anchors, key=lambda x: x[0])

    blocks: Dict[date, WeekBlock] = {}
    for i, (start_row, mon_date, _addr) in enumerate(anchors_sorted):
        end_row = ws_write.max_row if i == len(anchors_sorted) - 1 else anchors_sorted[i + 1][0] - 1
        blocks[mon_date] = WeekBlock(start_row, end_row, mon_date)

    for mon_date, block in blocks.items():
        in_status_conforming = False
        for r in range(block.anchor_row, block.end_row + 1):
            a1 = lower_clean(cell_value(ws_write, r, 1))
            b2 = normalize_str(cell_value(ws_write, r, 2))
            c3 = normalize_str(cell_value(ws_write, r, 3))
            b2_lower = b2.lower()
            c3_lower = c3.lower()

            # Status section starts at "status conforming" and ends at "comments" (in this template layout)
            if a1 == "status conforming" or b2_lower == "status conforming":
                in_status_conforming = True
            elif c3_lower == "comments" or a1 == "comments" or b2_lower == "comments":
                in_status_conforming = False

            if a1 == "owner" and block.owner_row is None:
                block.owner_row = r
            elif a1 == "property" and block.property_row is None:
                block.property_row = r
            elif a1 == "content type" and block.content_type_row is None:
                block.content_type_row = r
            elif a1 == "title" and block.title_row is None:
                block.title_row = r
            elif a1 == "synopsis" and block.synopsis_row is None:
                block.synopsis_row = r
            elif a1 == "planning strategy" and block.strategy_row is None:
                block.strategy_row = r
            elif a1 == "tx id / link" and block.tx_row is None:
                block.tx_row = r
            elif a1 == "tc in & out" and block.tc_row is None:
                block.tc_row = r
            # --- AÑADE ESTAS DOS LÍNEAS AQUÍ ---
            elif "burned" in a1 and "text" in a1 and block.burned_in_text_row is None:
                block.burned_in_text_row = r
            # -----------------------------------
            elif a1 == "comments" and block.comments_row is None:
                block.comments_row = r
            elif a1 == "comments" and block.comments_row is None:
                block.comments_row = r
            elif a1 == "deadline" and block.deadline_row is None:
                block.deadline_row = r

            if b2 in ("Tier A", "Tier B", "Tier C") and c3 in SHEET_COUNTRIES:
                if in_status_conforming:
                    block.status_country_rows.setdefault(c3, r)
                else:
                    block.rollout_country_rows.setdefault(c3, r)
            elif in_status_conforming and c3 in SHEET_COUNTRIES and b2 not in ("Tier A", "Tier B", "Tier C"):
                # In status conforming section, countries without Tier label
                block.status_country_rows.setdefault(c3, r)

        # Fallback partial match if some rows were not found
        needed = {
            "owner_row": "owner",
            "property_row": "property",
            "content_type_row": "content type",
            "title_row": "title",
            "synopsis_row": "synopsis",
            "strategy_row": "planning strategy",
            "tx_row": "tx id",
            "tc_row": "tc in & out",
            "burned_in_text_row": "burned",
            "comments_row": "comments",
            "deadline_row": "deadline",
        }
        if any(getattr(block, k) is None for k in needed):
            for r in range(block.anchor_row, block.end_row + 1):
                row_text = " ".join([
                    lower_clean(cell_value(ws_write, r, 1)),
                    lower_clean(cell_value(ws_write, r, 2)),
                    lower_clean(cell_value(ws_write, r, 3)),
                ])
                for attr, needle in needed.items():
                    if getattr(block, attr) is None and needle in row_text:
                        setattr(block, attr, r)

    return blocks

def infer_week_number_for_block(per_day_all_roles: Dict[str, List[Dict]], monday_date: date) -> str:
    """
    Infer 'Week Number (Year)' for the block from Airtable records.
    Use the most common non-empty value. Fallback to ISO week of monday_date.
    """
    values = []
    for recs in per_day_all_roles.values():
        for f in recs:
            wn = normalize_str(f.get(AT_FIELDS["week_number"]))
            if wn:
                values.append(wn)
    if values:
        most_common = Counter(values).most_common(1)[0][0]
        return most_common
    # fallback ISO week
    return str(monday_date.isocalendar()[1])