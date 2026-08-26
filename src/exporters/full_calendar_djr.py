import os
import logging
from collections import defaultdict, Counter
from datetime import timedelta, date

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from core.text_formatting import process_multiline_branding, markdown_to_richtext

from config.fields import AT_FIELDS
from config.layouts import DAY_ORDER
from config.layouts_djr import (
    DJR_TARGET_SHEET,
    DJR_DAY_LAYOUT,
    DJR_WEEK_LAYOUT,
)
from config.countries_djr import DJR_SHEET_COUNTRIES, DJR_REGION_PRESETS
from core.dates import (
    parse_date,
    month_name_from_records, 
    resolve_target_month_year, 
    first_visible_monday_of_target_month,
)

from core.text_formatting import normalize_str, lower_clean
from core.excel_utils import safe_set_value, center_wrap, cleanup_stale_merged_cells, get_writable_cell
from core.countries import collect_included_countries
from services.planner_service import (
    hub_role,
    asset_is_event,
    assemble_text_for_single_asset,
    earliest_deadline_for_day,
)
from exporters.monthly_overview_djr import populate_monthly_overview_djr
from exporters.shared.records import aggregate_by_visual_week_day_role  
from config.fields import AT_FIELDS
import warnings
import re

warnings.filterwarnings("ignore", message="wmf image format is not supported.*")
warnings.filterwarnings("ignore", message="Data Validation extension is not supported.*")

DC_TO_DJR_COUNTRY = {
    "France (French)": "France",
    "Germany": "GSA (German)",
    "Italy (Italian)": "Italy",
    "Middle East (Arabic - MSA)": "Middle East (Arabic)",
    "Poland (Polish)": "Poland",
    "Spain": "Spain",
    "Bulgaria (Bulgarian)": "Bulgaria",
    "Greece (Greek)": "Greece",
    "Israel (Hebrew)": "Israel (Hebrew)",
    "Netherlands (Dutch)": "Netherlands (Dutch)",
    "Portugal": "Portugal",
    "Romania (Romanian)": "Romania",
    "Turkey": "Türkiye",
    "Tukey": "Türkiye",
    "Disney Kids Turkey": "Türkiye",
    "Türkiye": "Türkiye",   
    "Sweden (Swedish)": "Sweden (Swedish)",
    "Denmark (Danish)": "Denmark (Danish)",
    "Norway (Norwegian)": "Norway",
}

def get_djr_target_columns(day_name, programming_count, mediaplanning_count):
    # Extraemos y ordenamos de izquierda a derecha (0, 1, 2, 3)
    all_cols = sorted([col for _, col in DJR_DAY_LAYOUT[day_name]])
    
    p_count = min(programming_count, 4)
    m_count = min(mediaplanning_count, 4 - p_count)
    
    # Asignación inteligente:
    # Programming prefiere los slots: 0, 1, 2, 3
    prog_slots = [0, 1, 2, 3][:p_count]
    prog_cols = [all_cols[i] for i in prog_slots]
    
    # Mediaplanning prefiere los slots: 2, 1, 0, 3 (Ignorando los que Programming ya tomó)
    remaining_slots = [i for i in [2, 1, 0, 3] if i not in prog_slots]
    media_slots = remaining_slots[:m_count]
    media_cols = [all_cols[i] for i in media_slots]
    
    used_slots = set(prog_slots + media_slots)
    
    uk_cols = []
    # El slot 3 (YT TEAM UK) se mantiene intacto si no hay 4 contenidos
    if 3 not in used_slots:
        uk_cols = [all_cols[3]]
        used_slots.add(3)
        
    extra_cols = [all_cols[i] for i in [0, 1, 2, 3] if i not in used_slots]
    
    return {"Programming": prog_cols, "Mediaplanning": media_cols, "UK": uk_cols, "Extra": extra_cols}

def write_djr_owner_row(ws, owner_row, day_name, column_map):
    for col in column_map["Programming"]: safe_set_value(ws, owner_row, col, "PROGRAMMING HUB", alignment=center_wrap())
    for col in column_map["Mediaplanning"]: safe_set_value(ws, owner_row, col, "MEDIAPLANNING HUB", alignment=center_wrap())
    for col in column_map["UK"]: safe_set_value(ws, owner_row, col, "YT TEAM UK", alignment=center_wrap())
    for col in column_map["Extra"]: safe_set_value(ws, owner_row, col, "", alignment=center_wrap())

def write_djr_start_date(ws, first_monday):
    cell = get_writable_cell(ws, 7, 4)
    cell.value = first_monday
    cell.alignment = center_wrap()
    cell.number_format = "dd-mmm-yy"

def write_djr_month_title(ws, records):
    cell_b1 = get_writable_cell(ws, 1, 2)
    cell_b1.value = "DISNEY JUNIOR YOUTUBE EMEA CALENDAR"
    cell_b1.alignment = center_wrap()

    target_year, target_month = resolve_target_month_year(records)
    if target_year and target_month:
        cell_b3 = get_writable_cell(ws, 3, 2)
        cell_b3.value = date(target_year, target_month, 1)
        cell_b3.alignment = center_wrap()
        cell_b3.number_format = "mmm-yy"

def write_djr_week_number(ws, layout, week_monday):
    week_num = week_monday.isocalendar()[1]
    safe_set_value(ws, layout["date_row"], 2, "WEEK", alignment=center_wrap())
    safe_set_value(ws, layout["date_row"], 3, week_num, alignment=center_wrap())

def write_djr_branding(ws, layout, col, record):
    branding_layout = layout.get("branding_rows")
    if not branding_layout: return
        
    f = record.get("fields", record)
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)

    def _set_rich_cell(row, column, text_val, alignment):
        if not row: return
        cell = ws.cell(row=row, column=column)
        cell.value = markdown_to_richtext(text_val)
        cell.alignment = alignment

    watermark_lines = process_multiline_branding(f.get(AT_FIELDS["branding_watermark"]), max_lines=4)
    for idx, row in enumerate(branding_layout["watermark"]):
        _set_rich_cell(row, col, watermark_lines[idx], align_center)

    bug_lines = process_multiline_branding(f.get(AT_FIELDS["branding_bug"]), max_lines=3)
    for idx, row in enumerate(branding_layout["bug"]):
        _set_rich_cell(row, col, bug_lines[idx], align_center)

    copyright_text = f.get(AT_FIELDS["branding_copyright"], "")
    _set_rich_cell(branding_layout["copyright"], col, copyright_text, align_left)

    endcard_lines = process_multiline_branding(f.get(AT_FIELDS["branding_endcard"]), max_lines=6)
    for idx, row in enumerate(branding_layout["endcard"]):
        _set_rich_cell(row, col, endcard_lines[idx], align_center)

def write_djr_asset_to_col(ws, layout, col, record):
    single = assemble_text_for_single_asset(record)
    content_type = single["content_type"]

    if normalize_str(record.get(AT_FIELDS["content_type"])) in {"Editorial Short", "Exclusive"}:
        content_type = "Exclusive Short"
    if normalize_str(record.get(AT_FIELDS["content_type"])) == "Vertical Short":
        content_type = "Vertical Short"

    safe_set_value(ws, layout["property_row"], col, single["property"], alignment=center_wrap())
    safe_set_value(ws, layout["content_type_row"], col, content_type, alignment=center_wrap())
    safe_set_value(ws, layout["title_row"], col, single["title"], alignment=center_wrap())
    safe_set_value(ws, layout["synopsis_row"], col, single["synopsis"], alignment=center_wrap())
    
    strat_val = record.get("fields", record).get(AT_FIELDS["strategy"])
    if isinstance(strat_val, list):
        actual_strategy = " / ".join(normalize_str(s) for s in strat_val if normalize_str(s))
    else:
        actual_strategy = normalize_str(strat_val)
    safe_set_value(ws, layout["strategy_row"], col, actual_strategy, alignment=center_wrap())
    
    safe_set_value(ws, layout["tx_row"], col, single["tx"], alignment=center_wrap())
    safe_set_value(ws, layout["tc_row"], col, single["tc"], alignment=center_wrap())
    write_djr_branding(ws, layout, col, record)

    role = hub_role(record.get(AT_FIELDS["owner_hub"]))
    comments_row = layout["comments_rows"].get(role)
    if comments_row:
        fields = record.get("fields", record)
        policy = fields.get("Rights Management Policy")
        policy_str = str(policy).strip() if policy else ""
        extra_note = ""
        
        if "Block AUS/NZ" in policy_str: extra_note = "\n\n⚠️ **Rights Policy:** Block AUS/NZ/TAIWAN/CHINA/HONG KONG - Monetize everywhere else"
        elif "Monetize Locally" in policy_str: extra_note = "\n\n⚠️ **Rights Policy:** Monetize Locally (Geo-block)"
        elif "Block Everywhere" in policy_str: extra_note = "\n\n🚫 **Rights Policy:** Block Everywhere / Do Not Monetize"
            
        comments_text = single["comments"] + extra_note
        safe_set_value(ws, comments_row, col, markdown_to_richtext(comments_text), alignment=center_wrap())

def write_djr_deadline(ws, layout, day_name, records):
    if not records: return
    deadline = earliest_deadline_for_day(records)
    if not deadline: return
    first_col = min(col for _, col in DJR_DAY_LAYOUT[day_name])
    cell = get_writable_cell(ws, layout["deadline_row"], first_col)
    cell.value = deadline
    cell.alignment = center_wrap()
    cell.number_format = "dd-mmm-yy"

def collect_included_countries_djr(record, channel_id_to_name_map=None):
    region_raw = record.get(AT_FIELDS["region"])
    region_str = str(region_raw).lower()
    
    if "spain only" in region_str:
        return {"Spain"}
        
    included = set()
    if region_raw:
        r_items = [lower_clean(x) for x in region_raw] if isinstance(region_raw, list) else [x.strip().lower() for x in str(region_raw).replace("[", "").replace("]", "").replace("'", "").replace('"', "").split(",") if x.strip()]
        for item in r_items:
            if item in DJR_REGION_PRESETS:
                included |= set(DJR_REGION_PRESETS[item])
                
    if not included:
        dc_included, _ = collect_included_countries(
            youtube_channel=record.get(AT_FIELDS["youtube_channel"]),
            region=region_raw,
            notes=record.get(AT_FIELDS["notes"]),
            excluded_markets=None,
            channel_id_to_name_map=channel_id_to_name_map,
        )
        included = {DC_TO_DJR_COUNTRY.get(c, c) for c in dc_included}
        
    ex_raw = record.get(AT_FIELDS["excluded_markets"])
    if ex_raw:
        ex_items = [lower_clean(x) for x in ex_raw] if isinstance(ex_raw, list) else [x.strip().lower() for x in str(ex_raw).split(",") if x.strip()]
        to_remove = set()
        for c in included:
            c_lower = c.lower() if isinstance(c, str) else ""
            for ex in ex_items:
                if ex in c_lower:
                    to_remove.add(c)
                for dc_k, dc_v in DC_TO_DJR_COUNTRY.items():
                    if dc_v == c and ex in dc_k.lower():
                        to_remove.add(c)
        included -= to_remove
        
    return included

def write_djr_country_rollout(ws, layout, column_map, record, col, channel_id_to_name_map):
    included = collect_included_countries_djr(record, channel_id_to_name_map)
    country_start, country_end = layout["country_rows"]
    status_start, status_end = layout["status_rows"]

    TIER_B_COUNTRIES = {"Poland", "Romania", "Netherlands (Dutch)", "Portugal", "Sweden (Swedish)", "Denmark (Danish)", "Norway", "Bulgaria"}
    FILL_TIER_A_C_YES = PatternFill(fill_type="solid", fgColor="EDEDED")
    FILL_TIER_B_YES = PatternFill(fill_type="solid", fgColor="FFFFFF")
    FILL_GREEN_YES = PatternFill(fill_type="solid", fgColor="00B050")

    for idx, country in enumerate(DJR_SHEET_COUNTRIES):
        rollout_row = country_start + idx
        status_row = status_start + idx
        value = "YES" if country in included else "NO"

        if value == "YES":
            font_rollout = Font(color="FFFFFFFF", bold=True)
            fill_rollout = FILL_GREEN_YES
        else:
            font_rollout = Font(color="FFFF0000", bold=True)
            fill_rollout = None 
        safe_set_value(ws, rollout_row, col, value, alignment=center_wrap(), font=font_rollout, fill=fill_rollout)

        if value == "YES":
            font_status = Font(color="00000000", bold=True)
            fill_status = FILL_TIER_B_YES if country in TIER_B_COUNTRIES else FILL_TIER_A_C_YES
        else:
            font_status = Font(color="FFFF0000", bold=True)
            fill_status = None 
        safe_set_value(ws, status_row, col, value, alignment=center_wrap(), font=font_status, fill=fill_status)

def rebuild_djr_week_merges(ws, layout):
    rows_to_rebuild = [layout["date_row"], layout["deadline_row"]]
    for cr in list(ws.merged_cells.ranges):
        if cr.min_row == cr.max_row and cr.min_row in rows_to_rebuild:
            ws.unmerge_cells(str(cr))

    for day_name in DAY_ORDER:
        cols = [col for _, col in DJR_DAY_LAYOUT[day_name]]
        start_col, end_col = min(cols), max(cols)
        ws.merge_cells(start_row=layout["date_row"], start_column=start_col, end_row=layout["date_row"], end_column=end_col)
        ws.merge_cells(start_row=layout["deadline_row"], start_column=start_col, end_row=layout["deadline_row"], end_column=end_col)

def shade_djr_out_of_month_days(ws, first_monday, target_year, target_month):
    fill = PatternFill(fill_type="solid", fgColor="D6DCE4")
    for week_index, layout in DJR_WEEK_LAYOUT.items():
        week_monday = first_monday + timedelta(days=(week_index - 1) * 7)
        for idx, day_name in enumerate(DAY_ORDER):
            current_date = week_monday + timedelta(days=idx)
            if current_date.year == target_year and current_date.month == target_month:
                continue
            cols = [col for _, col in DJR_DAY_LAYOUT[day_name]]
            for row in range(layout["owner_row"], layout["status_comments_row"] + 1):
                for col in cols:
                    ws.cell(row=row, column=col).fill = fill

def set_djr_column_widths(ws):
    for day_name in DAY_ORDER:
        for role, col in DJR_DAY_LAYOUT[day_name]:
            if role in {"Programming", "Mediaplanning"}:
                ws.column_dimensions[get_column_letter(col)].width = 35
            else:
                ws.column_dimensions[get_column_letter(col)].width = 10

def set_djr_row_heights(ws):
    for week_num, layout in DJR_WEEK_LAYOUT.items():
        # Filas estandar a 30
        for row in [
            layout["property_row"], layout["content_type_row"], layout["title_row"],
            layout["synopsis_row"], layout["strategy_row"], layout["tx_row"], layout["tc_row"]
        ]:
            ws.row_dimensions[row].height = 30
            
        # Filas de branding y copyright (La más alta)
        branding_layout = layout.get("branding_rows")
        if branding_layout:
            for br_row in branding_layout.get("watermark", []) + branding_layout.get("bug", []) + branding_layout.get("endcard", []):
                if br_row: ws.row_dimensions[br_row].height = 38
            
            # Fila exclusiva de Copyright (40, 102, 164, etc.)
            copy_row = branding_layout.get("copyright")
            if copy_row: ws.row_dimensions[copy_row].height = 95
                
        # Filas de comments a 35 (Status Conforming ya no se toca)
        for c_row in layout["comments_rows"].values():
            if c_row: ws.row_dimensions[c_row].height = 35


# --- NUEVO: Exportar ID's (Overon Tracking) ---
def export_overon_ids(wb, records):
    sheet_name = "ID's"
    if sheet_name not in wb.sheetnames:
        return

    ws_ids = wb[sheet_name]
    current_row = 4

    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    thin_border = Border(left=Side(style='thin', color="000000"), right=Side(style='thin', color="000000"), top=Side(style='thin', color="000000"), bottom=Side(style='thin', color="000000"))
    header_fill = PatternFill(fill_type="solid", fgColor="F8CBAD")

    def _write_cell(ws, r, c, val, align=align_center, is_header=False):
        cell = ws.cell(row=r, column=c)
        cell.value = val
        cell.border = thin_border
        cell.alignment = align
        if is_header:
            cell.fill = header_fill
            cell.font = Font(bold=True)

    def _apply_merged_border(ws, r_start, c_start, r_end, c_end):
        for r in range(r_start, r_end + 1):
            for c in range(c_start, c_end + 1):
                ws.cell(row=r, column=c).border = thin_border

    valid_records = []
    for r in records:
        f = r.get("fields", r)
        owner = hub_role(f.get(AT_FIELDS["owner_hub"]))
        ctype = normalize_str(f.get(AT_FIELDS["content_type"])).lower()
        subtype = normalize_str(f.get("Content Subtype", "")).lower()
        title = normalize_str(f.get(AT_FIELDS["youtube_title"], "")).lower()
        
        is_prog = (owner == "Programming")
        is_comp = ("compilation" in ctype)
        is_mega = ("mega" in ctype or "mega" in subtype or "mega" in title)
        is_live = ("live stream" in ctype or "live" in subtype)
        
        tx_raw = normalize_str(f.get(AT_FIELDS["tx_id"]))
        tcin_raw = normalize_str(f.get(AT_FIELDS["tc_in"]))
        tcout_raw = normalize_str(f.get(AT_FIELDS["tc_out"]))
        has_tcs = bool(tx_raw and tcin_raw and tcout_raw)
        
        valid = False
        if is_prog:
            if is_comp and has_tcs: valid = True
            if is_mega: valid = True
            if is_live and is_mega: valid = True
            
        if valid:
            live_date_str = f.get(AT_FIELDS["live_date"])
            parsed_date = parse_date(live_date_str) if live_date_str else date.min
            valid_records.append((parsed_date, f))
            
    valid_records.sort(key=lambda x: x[0])

    for _, f in valid_records:
        activity_name = normalize_str(f.get("Activity") or f.get("Activity Name") or f.get("Asset Label") or f.get(AT_FIELDS["youtube_title"]))
        live_date = parse_date(f.get(AT_FIELDS["live_date"]))
        date_str = f"{live_date.day:02d}/{live_date.month:02d}" if live_date else ""
        
        tx_raw = normalize_str(f.get(AT_FIELDS["tx_id"]))
        tcin_raw = normalize_str(f.get(AT_FIELDS["tc_in"]))
        tcout_raw = normalize_str(f.get(AT_FIELDS["tc_out"]))
        cd_raw = normalize_str(f.get("Content Details", ""))
        notes = normalize_str(f.get(AT_FIELDS["notes"]))
        
        tx_lines = [x.strip() for x in tx_raw.split('\n') if x.strip()]
        tcin_lines = [x.strip() for x in tcin_raw.split('\n') if x.strip()]
        tcout_lines = [x.strip() for x in tcout_raw.split('\n') if x.strip()]
        cd_lines = [x.strip() for x in cd_raw.split('\n') if x.strip()]
        
        ctype = normalize_str(f.get(AT_FIELDS["content_type"])).lower()
        subtype = normalize_str(f.get("Content Subtype", "")).lower()
        title = normalize_str(f.get(AT_FIELDS["youtube_title"], "")).lower()
        is_mega = ("mega" in ctype or "mega" in subtype or "mega" in title)
        has_cd = bool(cd_raw)
        
        parsed_items = []
        if is_mega and not (has_cd and len(cd_lines) > 1):
            sub_titles_from_notes = [line.strip() for line in notes.split('\n') if line.strip() and "compile" not in line.lower() and "reuse" not in line.lower()]
            current_sub = sub_titles_from_notes[0] if sub_titles_from_notes else "Part 1"
            tc_idx = 0
            
            if tx_lines:
                for tx_val in tx_lines:
                    if "INTL" not in tx_val.upper() and "ILJ" not in tx_val.upper() and "http" not in tx_val.lower() and len(tx_val) > 12:
                        current_sub = tx_val
                    else:
                        tcin_val = tcin_lines[tc_idx] if tc_idx < len(tcin_lines) else ""
                        tcout_val = tcout_lines[tc_idx] if tc_idx < len(tcout_lines) else ""
                        parsed_items.append({"sub": current_sub, "tx": tx_val, "in": tcin_val, "out": tcout_val})
                        tc_idx += 1
            else:
                if sub_titles_from_notes:
                    for sub in sub_titles_from_notes: parsed_items.append({"sub": sub, "tx": "", "in": "", "out": ""})
                else:
                    parsed_items.append({"sub": notes, "tx": "", "in": "", "out": ""})
        else:
            num_items = max(len(tx_lines), len(tcin_lines), len(tcout_lines), len(cd_lines), 1)
            for i in range(num_items):
                tx_val = tx_lines[i] if i < len(tx_lines) else ""
                tcin_val = tcin_lines[i] if i < len(tcin_lines) else ""
                tcout_val = tcout_lines[i] if i < len(tcout_lines) else ""
                cd_val = cd_lines[i] if i < len(cd_lines) else ""
                parsed_items.append({"sub": cd_val, "tx": tx_val, "in": tcin_val, "out": tcout_val})
        
        is_4_cols = is_mega or has_cd
        headers = ["Content Details", "ID", "TC IN", "TC OUT"] if is_4_cols else ["ID", "TC IN", "TC OUT"]
        start_merge_a = current_row + 1
        
        for i, h in enumerate(headers): _write_cell(ws_ids, current_row, 2 + i, h, is_header=True)
        _write_cell(ws_ids, current_row, 1, "", is_header=True)
        
        current_row += 1
        
        if is_4_cols:
            grouped = defaultdict(list)
            for item in parsed_items: grouped[item["sub"]].append(item)
                
            for sub, items in grouped.items():
                start_b = current_row
                for item in items:
                    _write_cell(ws_ids, current_row, 1, "")
                    _write_cell(ws_ids, current_row, 2, item["sub"])
                    _write_cell(ws_ids, current_row, 3, item["tx"], align_left)
                    _write_cell(ws_ids, current_row, 4, item["in"])
                    _write_cell(ws_ids, current_row, 5, item["out"])
                    current_row += 1
                
                if current_row - 1 > start_b:
                    ws_ids.merge_cells(start_row=start_b, start_column=2, end_row=current_row-1, end_column=2)
                    _apply_merged_border(ws_ids, start_b, 2, current_row-1, 2)
                    ws_ids.cell(row=start_b, column=2).alignment = align_center
        else:
            for item in parsed_items:
                _write_cell(ws_ids, current_row, 1, "")
                _write_cell(ws_ids, current_row, 2, item["tx"], align_left)
                _write_cell(ws_ids, current_row, 3, item["in"])
                _write_cell(ws_ids, current_row, 4, item["out"])
                current_row += 1
                
        if current_row - 1 >= start_merge_a:
            if current_row - 1 > start_merge_a:
                ws_ids.merge_cells(start_row=start_merge_a, start_column=1, end_row=current_row-1, end_column=1)
            _apply_merged_border(ws_ids, start_merge_a, 1, current_row-1, 1)
            c_a = ws_ids.cell(row=start_merge_a, column=1)
            c_a.value = f"{activity_name} {date_str}"
            c_a.alignment = align_center
            
        current_row += 1 

# --- NUEVO: Exportar Metadata Programming Spain ---
def export_metadata_spain(wb, records, first_monday):
    sheet_name = "Metadata Programming Spain"
    if sheet_name not in wb.sheetnames:
        return

    ws_meta = wb[sheet_name]
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    font_green = Font(color="00B050", bold=True)
    fill_grey = PatternFill(fill_type="solid", fgColor="D6DCE4")
    fill_white = PatternFill(fill_type="solid", fgColor="FFFFFF")

    cell_b1 = ws_meta.cell(row=1, column=2)
    cell_b1.value = "DISNEY JUNIOR YOUTUBE EMEA CALENDAR"
    cell_b1.alignment = align_center

    target_year, target_month = resolve_target_month_year(records)
    if target_year and target_month:
        cell_b3 = ws_meta.cell(row=3, column=2)
        cell_b3.value = date(target_year, target_month, 1)
        cell_b3.alignment = align_center
        cell_b3.number_format = "mmm-yy"

    # Inyección a prueba de fallos en la celda combinada D7:G7 (Escribimos sin invocar "merge" manual)
    cell_d7 = ws_meta.cell(row=7, column=4)
    cell_d7.value = first_monday
    cell_d7.alignment = align_center
    cell_d7.number_format = "dd-mmm-yy"

    META_WEEKS = {1: 6, 2: 16, 3: 26, 4: 36, 5: 46}
    prog_records = [r for r in records if hub_role(r.get("fields", r).get(AT_FIELDS["owner_hub"])) == "Programming"]
    buckets = aggregate_by_visual_week_day_role(prog_records, first_monday)

    for week_index, start_row in META_WEEKS.items():
        week_monday = first_monday + timedelta(days=(week_index - 1) * 7)
        ws_meta.cell(row=start_row, column=3).value = week_monday.isocalendar()[1]
        
        for idx, day_name in enumerate(DAY_ORDER):
            current_date = week_monday + timedelta(days=idx)
            day_cols = [col for role, col in DJR_DAY_LAYOUT[day_name] if role == "Programming"]
            if not day_cols: continue
            
            start_col = min([col for _, col in DJR_DAY_LAYOUT[day_name]])
            end_col = max([col for _, col in DJR_DAY_LAYOUT[day_name]])
            
            is_target_month = (current_date.month == target_month)
            active_fill = fill_white if is_target_month else fill_grey
            
            # Pintamos el fondo para TODO el bloque (desde start_col hasta end_col)
            for r in range(start_row + 3, start_row + 8):
                for c in range(start_col, end_col + 1):
                    ws_meta.cell(row=r, column=c).fill = active_fill
            
            day_records = buckets.get((week_index, day_name, "Programming"), [])

            for i, record in enumerate(day_records[:len(day_cols)]):
                col = day_cols[i]
                f = record.get("fields", record)
                single = assemble_text_for_single_asset(record)
                
                es_title = normalize_str(f.get("ES Title", single["title"]))
                es_desc = normalize_str(f.get("ES Description", single["synopsis"]))

                safe_set_value(ws_meta, start_row + 2, col, "PROGRAMMING HUB", alignment=align_center)
                safe_set_value(ws_meta, start_row + 3, col, single["property"], alignment=align_center)
                safe_set_value(ws_meta, start_row + 4, col, single["content_type"], alignment=align_center)
                safe_set_value(ws_meta, start_row + 5, col, es_title, alignment=align_center, font=font_green)
                safe_set_value(ws_meta, start_row + 6, col, es_desc, alignment=align_center, font=font_green)
                safe_set_value(ws_meta, start_row + 7, col, single["strategy"], alignment=align_center)

def populate_calendar_djr(
    template_path: str,
    output_path: str,
    records: list,
    enable_mediaplanning: bool = True,
    channel_id_to_name_map: dict | None = None,
):
    print("ENTERED populate_calendar_djr")
    if channel_id_to_name_map is None: channel_id_to_name_map = {}

    wb = load_workbook(template_path, data_only=False)
    if DJR_TARGET_SHEET not in wb.sheetnames: raise RuntimeError(f"Target sheet '{DJR_TARGET_SHEET}' not found in workbook.")
    ws = wb[DJR_TARGET_SHEET]

    set_djr_column_widths(ws)
    write_djr_month_title(ws, records)

    first_monday = first_visible_monday_of_target_month(records)
    if not first_monday: return

    write_djr_start_date(ws, first_monday)
    buckets = aggregate_by_visual_week_day_role(records, first_monday)

    for week_index, layout in DJR_WEEK_LAYOUT.items():
        week_monday = first_monday + timedelta(days=(week_index - 1) * 7)
        write_djr_week_number(ws, layout, week_monday)

        for day_name in DAY_ORDER:
            prog_records = buckets.get((week_index, day_name, "Programming"), [])
            media_records = buckets.get((week_index, day_name, "Mediaplanning"), [])
            if not enable_mediaplanning: media_records = []

            column_map = get_djr_target_columns(day_name, len(prog_records), len(media_records))
            write_djr_owner_row(ws, layout["owner_row"], day_name, column_map)
            write_djr_deadline(ws, layout, day_name, prog_records + media_records)

            for idx, record in enumerate(prog_records[:len(column_map["Programming"])]):
                col = column_map["Programming"][idx]
                write_djr_asset_to_col(ws, layout, col, record)
                write_djr_country_rollout(ws, layout, column_map, record, col, channel_id_to_name_map)

            for idx, record in enumerate(media_records[:len(column_map["Mediaplanning"])]):
                col = column_map["Mediaplanning"][idx]
                write_djr_asset_to_col(ws, layout, col, record)
                write_djr_country_rollout(ws, layout, column_map, record, col, channel_id_to_name_map)

        rebuild_djr_week_merges(ws, layout)

    target_year, target_month = resolve_target_month_year(records)
    if first_monday and target_year and target_month:
        shade_djr_out_of_month_days(ws, first_monday, target_year, target_month)

    set_djr_row_heights(ws)

    export_overon_ids(wb, records)
    export_metadata_spain(wb, records, first_monday)

    try:
        populate_monthly_overview_djr(wb=wb, records=records, enable_mediaplanning=enable_mediaplanning, channel_id_to_name_map=channel_id_to_name_map)
    except Exception as e:
        pass

    out_dir = os.path.dirname(output_path)
    if out_dir: os.makedirs(out_dir, exist_ok=True)

    cleanup_stale_merged_cells(ws)
    wb.template = False

    month_name_upper = month_name_from_records(records).capitalize()
    month_abbr = month_name_upper[:3] 
    for sn in wb.sheetnames:
        if "highlights" in sn.lower():
            new_title = f"YT DJR EMEA {month_abbr} Highlights"
            wb[sn].title = new_title
            if target_year and target_month:
                ws_high = wb[new_title]
                safe_set_value(ws_high, 3, 5, date(target_year, target_month, 1), number_format="mmm-yy")

    cleanup_stale_merged_cells(ws) 
    wb.save(output_path)