import logging
import math
from collections import defaultdict, Counter
from datetime import date, timedelta
from openpyxl.styles import Font, Alignment, PatternFill
from typing import Dict, List, Optional, Tuple
from openpyxl.styles import Font, Alignment, PatternFill
from services.planner_service import (
    hub_role,
    asset_is_event,
)
from config.content_types import MONTHLY_OVERVIEW_MAPPING
from config.fields import AT_FIELDS
from config.layouts import(
    DAY_ORDER,
    OVERVIEW_DAY_COLS,
    OVERVIEW_WEEK_LAYOUT,
    MONTHLY_OVERVIEW_SHEET,
)
from config.styles import OVERVIEW_FILL_COLORS, OVERVIEW_DAY_LABELS
from core.excel_utils import get_writable_cell

from core.countries import (
    collect_included_countries, 

    normalize_str,
    lower_clean,
    normalize_overview_region,
    format_except_suffix
)

from core.dates import (
    parse_date, 
    first_visible_monday,
    month_name_from_records,
)

from core.text_formatting import replace_month_in_title

def overview_bucket(priority: str) -> str:
    p = normalize_str(priority)
    if p in {"Current A", "Current B", "Current C"}:
        return "priority"
    if p in {"Tactical A", "Tactical B", "Tactical Opportunity"}:
        return "tactical"
    if p == "Evergreen":
        return "evergreen"
    return "priority"

def overview_row_offset(content_type: str, priority: str) -> Optional[int]:
    ct = normalize_str(content_type)
    bucket = overview_bucket(priority)
    
    # 1. Búsqueda a prueba de balas (case-insensitive)
    mapped_category = ct
    for key, value in MONTHLY_OVERVIEW_MAPPING.items():
        if key.lower() == ct.lower():
            mapped_category = value
            break

    # Convertimos mapped_category y ct a minúsculas para evaluar
    cat_low = mapped_category.lower()
    ct_low = ct.lower()

    # Full Eps
    if cat_low == "full episode":
        return {
            "priority": 0, "tactical": 1, "evergreen": 2, "acquisitions": 3
        }.get(bucket, 0)

    # Exclusive Shorts
    if cat_low == "exclusive":
        return 4

    # Scene Lifts / Sneak Peeks / Videoclips
    if cat_low == "scene lifts / sneak peeks / videoclips" or ct_low in {"lift", "musical lift", "sneak peek", "videoclip"}:
        return 5

    # Special Formats (Fillers)
    if cat_low in {"podcast", "magical moment", "interactive", "editorial short"}:
        return 6

    # Compilations
    if cat_low == "compilation":
        return {
            "priority": 7, "tactical": 8, "evergreen": 9, "acquisitions": 10
        }.get(bucket, 7)

    # Live Streams
    if cat_low == "live stream":
        return {
            "priority": 11, "tactical": 12, "evergreen": 13, "acquisitions": 14
        }.get(bucket, 11)

    # Vertical Shorts
    if cat_low == "vertical short":
        return 15

    return None

def overview_visual_position(live_date: date, first_monday: date) -> Optional[Tuple[int, str]]:
    """
    Returns (visual_week_number, day_name) based on the visible calendar grid.
    Week numbers are 1..5 relative to first visible Monday.
    """
    if not live_date or not first_monday:
        return None

    delta_days = (live_date - first_monday).days
    if delta_days < 0:
        return None

    visual_week = (delta_days // 7) + 1
    if visual_week < 1 or visual_week > 5:
        return None

    day_name = DAY_ORDER[live_date.weekday()]
    return visual_week, day_name

def build_overview_lines(record: Dict, channel_id_to_name_map: Optional[Dict[str, str]] = None) -> Tuple[str, str, bool]:
    """
    Returns:
    line1, line2, geofenced
    """
    if channel_id_to_name_map is None:
        channel_id_to_name_map = {}

    activity = normalize_str(record.get(AT_FIELDS["activity"]))
    if not activity:
        activity = normalize_str(record.get(AT_FIELDS["asset_label"])) or normalize_str(record.get(AT_FIELDS["title"]))

    content_type = normalize_str(record.get(AT_FIELDS["content_type"]))

    if content_type == "Musical Lift":
        activity = f"♪ {activity}"
    elif content_type == "Lift":
        activity = f"▶ {activity}"
    elif content_type == "Exclusive":
        activity = f"★ {activity}"
    elif content_type == "Compilation":
        activity = f"COMP {activity}"
    elif content_type == "Live Stream":
        activity = f"LIVE {activity}"

    geofenced = lower_clean(record.get(AT_FIELDS["geofenced"])) == "yes"

    included_countries, _ = collect_included_countries(
        youtube_channel=record.get(AT_FIELDS["youtube_channel"]),
        region=record.get(AT_FIELDS["region"]),
        notes=record.get(AT_FIELDS["notes"]),
        excluded_markets=record.get(AT_FIELDS["excluded_markets"]),
        channel_id_to_name_map=channel_id_to_name_map,
    )

    line2 = normalize_overview_region(record.get(AT_FIELDS["region"]))
    except_suffix = format_except_suffix(record.get(AT_FIELDS["region"]), included_countries)

    if line2:
        line2 = f"{line2}{except_suffix}"
    else:
        line2 = except_suffix.strip()

    return activity, line2, geofenced

def apply_overview_fill(cell, content_type: str):
    fill_color = OVERVIEW_FILL_COLORS.get(normalize_str(content_type))
    if fill_color:
        cell.fill = PatternFill(fill_type="solid", fgColor=fill_color)

def write_overview_cell(ws, row: int, col: int, line1: str, line2: str, content_type: str, geofenced: bool):
    cell = get_writable_cell(ws, row, col)

    new_text = line1 if not line2 else f"{line1}\n{line2}"

    existing = normalize_str(cell.value)
    if existing:
        cell.value = f"{existing}\n\n— — —\n\n{new_text}"
    else:
        cell.value = new_text

    # --- RESOLVER COLOR CON MAPEO CENTRALIZADO (ALIAS SAFE) ---
    ct_norm = normalize_str(content_type)
    mapped_category = ct_norm
    for key, value in MONTHLY_OVERVIEW_MAPPING.items():
        if key.lower() == ct_norm.lower():
            mapped_category = value
            break

    fill_color = OVERVIEW_FILL_COLORS.get(ct_norm)
    if not fill_color:
        # Si pertenece al bucket de Lifts/Sneak Peeks/Videoclips, forzar el color verde solicitado #A9D08E
        if mapped_category.lower() == "scene lifts / sneak peeks / videoclips" or ct_norm.lower() in {"lift", "musical lift", "sneak peek", "videoclip"}:
            fill_color = "FFA9D08E"
        else:
            fill_color = OVERVIEW_FILL_COLORS.get(mapped_category)

    if fill_color:
        cell.fill = PatternFill(fill_type="solid", fgColor=fill_color)

    font_color = "FF800080" if geofenced else "FF000000"
    cell.font = Font(bold=True, color=font_color)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

def write_vertical_short_cell(ws, top_row: int, col: int, record: Dict):
    """
    Vertical Shorts use 2 stacked cells:
    - top cell: PROPERTY (uppercase) + Strategy, centered, black
    - bottom cell: Region/Tiers, centered, bold, red
      or yellow if SPAIN ONLY
    Ignores geofenced purple styling for this content type.
    """
    top_cell = get_writable_cell(ws, top_row, col)
    bottom_cell = get_writable_cell(ws, top_row + 1, col)

    prop_val = record.get(AT_FIELDS["property"])
    if isinstance(prop_val, list):
        prop = " / ".join(normalize_str(p) for p in prop_val if normalize_str(p))
    else:
        prop = normalize_str(prop_val)
    prop = prop.upper()

    strat_val = record.get(AT_FIELDS["strategy"])
    if isinstance(strat_val, list):
        strat = " / ".join(normalize_str(s) for s in strat_val if normalize_str(s))
    else:
        strat = normalize_str(strat_val)

    tier_text = normalize_overview_region(record.get(AT_FIELDS["region"]))

    # top cell text
    top_text = prop
    if strat:
        top_text = f"{prop}\n{strat}"

    top_cell.value = top_text
    top_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    top_cell.font = Font(bold=True, color="FF000000")
    top_cell.fill = PatternFill(fill_type="solid", fgColor=OVERVIEW_FILL_COLORS["Vertical Short"])

    # bottom cell text
    bottom_cell.value = tier_text
    bottom_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    if tier_text == "SPAIN ONLY":
        bottom_cell.font = Font(bold=True, color="FFFFFF00")  # yellow
    else:
        bottom_cell.font = Font(bold=True, color="FFFF0000")  # red

    bottom_cell.fill = PatternFill(fill_type="solid", fgColor=OVERVIEW_FILL_COLORS["Vertical Short"])

def auto_adjust_overview_row_height(ws, row: int, min_height: float = 24, per_line: float = 16, chars_per_line: int = 22):
    """
    Better approximation of row auto-height based on explicit line breaks and wrapped text.
    """
    max_visual_lines = 1

    for col in range(4, 11):  # D:J
        cell = get_writable_cell(ws, row, col)
        value = "" if cell.value is None else str(cell.value)

        if not value:
            continue

        visual_lines = 0
        for line in value.split("\n"):
            line = line.strip()
            if not line:
                visual_lines += 1
            else:
                visual_lines += max(1, math.ceil(len(line) / chars_per_line))

        max_visual_lines = max(max_visual_lines, visual_lines)

    ws.row_dimensions[row].height = max(min_height, max_visual_lines * per_line)

def clear_overview_week(ws, week_num: int):
    """
    Clear only content/event writable area in a week block.
    Do NOT clear header/date formula cells.
    """
    if week_num not in OVERVIEW_WEEK_LAYOUT:
        return

    layout = OVERVIEW_WEEK_LAYOUT[week_num]

    # Clear only event rows
    for row in layout["event_rows"]:
        for col in range(4, 11):
            cell = get_writable_cell(ws, row, col)
            cell.value = None

    # Clear only content area
    for row in range(layout["base_row"], layout["end_row"] + 1):
        for col in range(4, 11):
            cell = get_writable_cell(ws, row, col)
            cell.value = None

def write_overview_headers(ws, records: List[Dict]):
    month_name_upper = month_name_from_records(records)

    title_cell = get_writable_cell(ws, 2, 2)  # B2
    existing_title = normalize_str(title_cell.value)
    if existing_title:
        title_cell.value = replace_month_in_title(existing_title, month_name_upper)
    else:
        title_cell.value = f"DC EMEA YT {month_name_upper} 2026 CALENDAR"
    title_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for week_num, layout in OVERVIEW_WEEK_LAYOUT.items():
        for day_name, col in OVERVIEW_DAY_COLS.items():
            day_label_cell = get_writable_cell(ws, layout["header_row"], col)
            day_label_cell.value = OVERVIEW_DAY_LABELS[day_name]
            day_label_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    first_monday = first_visible_monday(records)
    if first_monday:
        monday_date_cell = get_writable_cell(ws, 5, OVERVIEW_DAY_COLS["Monday"])  # D5 only
        monday_date_cell.value = first_monday
        monday_date_cell.number_format = "dd-mmm"
        monday_date_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def collect_overview_events(records: List[Dict]) -> Dict[int, List[Tuple[str, int, int]]]:
    """
    Returns:
    { week_num: [(event_text, start_col, end_col), ...] }

    Uses EVENT content type and merges across consecutive days if same text.
    """
    events_by_week = defaultdict(list)

    raw = []
    for rec in records:
        fields = rec.get("fields", rec)
        if normalize_str(fields.get(AT_FIELDS["content_type"])) != "EVENT":
            continue

        live = parse_date(fields.get(AT_FIELDS["live_date"]))
        if not live:
            continue

        wom = fields.get(AT_FIELDS["week_of_month"])
        try:
            week_num = int(str(wom).strip())
        except Exception:
            continue

        day_name = normalize_str(fields.get(AT_FIELDS["day_of_week"]))
        if day_name not in DAY_ORDER:
            day_name = DAY_ORDER[live.weekday()]

        event_text = normalize_str(fields.get(AT_FIELDS["activity"])) or normalize_str(fields.get(AT_FIELDS["title"]))
        if not event_text:
            continue

        raw.append((week_num, day_name, event_text))

    # group by week + text
    grouped = defaultdict(list)
    for week_num, day_name, event_text in raw:
        grouped[(week_num, event_text)].append(OVERVIEW_DAY_COLS[day_name])

    for (week_num, event_text), cols in grouped.items():
        cols = sorted(set(cols))
        # merge contiguous spans
        start = cols[0]
        prev = cols[0]
        for col in cols[1:]:
            if col == prev + 1:
                prev = col
            else:
                events_by_week[week_num].append((event_text, start, prev))
                start = col
                prev = col
        events_by_week[week_num].append((event_text, start, prev))

    return events_by_week

def write_overview_events(ws, records: List[Dict]):
    """
    Writes EVENT labels into grey rows.
    Uses 2 rows max per week.
    If event spans multiple days, merges across those day columns.
    """
    events_by_week = collect_overview_events(records)

    for week_num, event_spans in events_by_week.items():
        if week_num not in OVERVIEW_WEEK_LAYOUT:
            continue

        event_rows = OVERVIEW_WEEK_LAYOUT[week_num]["event_rows"]

        for idx, (event_text, start_col, end_col) in enumerate(event_spans[:2]):
            row = event_rows[idx]

            # merge if span > 1 col
            if end_col > start_col:
                ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)

            cell = get_writable_cell(ws, row, start_col)
            cell.value = event_text
            # CAMBIO: color="FFFFFFFF" (Blanco) y bold=True para legibilidad
            cell.font = Font(italic=True, bold=True, color="FFFFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for week_num, layout in OVERVIEW_WEEK_LAYOUT.items():
        for event_row in layout["event_rows"]:
            auto_adjust_overview_row_height(ws, event_row, min_height=18, per_line=14)

def populate_monthly_overview(
    wb,
    records: List[Dict],
    enable_mediaplanning: bool = True,
    channel_id_to_name_map: Optional[Dict[str, str]] = None,
):
    """
    Populate 'MONTHLY OVERVIEW YT DC EMEA'.

    Uses visual calendar positioning based on the first visible Monday,
    not Airtable Week of Month.
    """
    if channel_id_to_name_map is None:
        channel_id_to_name_map = {}    
    
    if MONTHLY_OVERVIEW_SHEET not in wb.sheetnames:
        logging.warning(f"Sheet '{MONTHLY_OVERVIEW_SHEET}' not found. Skipping MONTHLY OVERVIEW export.")
        return

    ws = wb[MONTHLY_OVERVIEW_SHEET]

    # Header/title and first Monday
    write_overview_headers(ws, records)

    # Events
    write_overview_events(ws, records)

    first_monday = first_visible_monday(records)
    if not first_monday:
        logging.warning("Could not determine first visible Monday for MONTHLY OVERVIEW.")
        return

    used_weeks = set()

    for rec in records:
        fields = rec.get("fields", rec)

        if asset_is_event(fields):
            continue

        owner = hub_role(fields.get(AT_FIELDS["owner_hub"]))
        if owner not in {"Programming", "Mediaplanning"}:
            continue

        if owner == "Mediaplanning" and not enable_mediaplanning:
            continue

        content_type = normalize_str(fields.get(AT_FIELDS["content_type"]))
        live_date = parse_date(fields.get(AT_FIELDS["live_date"]))
        if not live_date:
            continue

        pos = overview_visual_position(live_date, first_monday)
        if not pos:
            continue

        visual_week, day_name = pos
        if visual_week not in OVERVIEW_WEEK_LAYOUT:
            continue

        used_weeks.add(visual_week)

        row_offset = overview_row_offset(content_type, fields.get(AT_FIELDS["priority"]))
        if row_offset is None:
            continue

        base_row = OVERVIEW_WEEK_LAYOUT[visual_week]["base_row"]
        target_row = base_row + row_offset
        target_col = OVERVIEW_DAY_COLS[day_name]

        if content_type == "Vertical Short":
            write_vertical_short_cell(ws, target_row, target_col, fields)
            auto_adjust_overview_row_height(ws, target_row, min_height=26, per_line=16)
            auto_adjust_overview_row_height(ws, target_row + 1, min_height=22, per_line=16)
        else:
            line1, line2, geofenced = build_overview_lines(fields, channel_id_to_name_map=channel_id_to_name_map)
            write_overview_cell(ws, target_row, target_col, line1, line2, content_type, geofenced)
            auto_adjust_overview_row_height(ws, target_row)

    # If week 5 unused, clear it
    if 5 not in used_weeks:
        clear_overview_week(ws, 5)

    # --- NEW: Out-of-month Reverse Diagonal Stripe Pattern ---
    # 1. Determinar cuál es el mes principal ("Target Month")
    valid_dates = [parse_date(r.get("fields", r).get(AT_FIELDS["live_date"])) for r in records]
    valid_dates = [d for d in valid_dates if d]
    
    if valid_dates and first_monday:
        target_month = Counter(d.month for d in valid_dates).most_common(1)[0][0]

        # En openpyxl: lightDown = Reverse Diagonal. 
        # fgColor = color del patrón (Líneas negras). bgColor = color de fondo (Gris D9D9D9).
        out_of_month_fill = PatternFill(fill_type="lightDown", fgColor="FF000000", bgColor="FFD9D9D9")
        days_list = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]

        for week_num, layout in OVERVIEW_WEEK_LAYOUT.items():
            base_row = layout["base_row"]
            end_row = layout["end_row"]

            for day_idx, day_name in enumerate(days_list):
                col = OVERVIEW_DAY_COLS.get(day_name)
                if not col:
                    continue

                # Calcular la fecha exacta de esta columna en el layout
                current_date = first_monday + timedelta(days=(week_num - 1) * 7 + day_idx)

                # Si el día no pertenece al mes principal, aplicamos el patrón a las filas de contenido vacías
                if current_date.month != target_month:
                    for row in range(base_row, end_row + 1):
                        cell = get_writable_cell(ws, row, col)
                        # Si la celda está vacía, aplicamos el patrón gris sin importar el formato de la plantilla
                        if not cell.value:
                            cell.fill = out_of_month_fill

    # Event row height approximation
    for week_num, layout in OVERVIEW_WEEK_LAYOUT.items():
        for event_row in layout["event_rows"]:
            auto_adjust_overview_row_height(ws, event_row, min_height=24, per_line=14)

    logging.info(f"Finished populating '{MONTHLY_OVERVIEW_SHEET}'.")
