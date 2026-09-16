import logging
import re
from collections import defaultdict
from datetime import timedelta

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import Cell, MergedCell

from config.fields import AT_FIELDS
from config.layouts import DAY_ORDER
from config.layouts_djr import DJR_MONTHLY_OVERVIEW_SHEET
from config.countries_djr import DJR_COUNTRY_SHORT_NAMES
from core.dates import (
    parse_date,
    first_visible_monday_of_target_month,
    month_name_from_records,
)
from core.text_formatting import normalize_str, lower_clean
from core.excel_utils import get_writable_cell, center_wrap
from core.countries import normalize_overview_region
from services.planner_service import hub_role, asset_is_event
from typing import Dict, List, Optional, Tuple

DJR_OVERVIEW_DAY_COLS = {
    "Monday": 4, "Tuesday": 5, "Wednesday": 6, "Thursday": 7,
    "Friday": 8, "Saturday": 9, "Sunday": 10,
}

# --- ACTUALIZADO: Layout con filas separadas para Interactive y Podcast ---
DJR_OVERVIEW_WEEK_LAYOUT = {
    1: {"date_row": 9, "full_eps": 10, "exclusive": 11, "lifts": 12, "compilations": 13, "live": 14, "interactive": 15, "podcast": 16, "vertical_rows": (17, 18)},
    2: {"date_row": 20, "full_eps": 21, "exclusive": 22, "lifts": 23, "compilations": 24, "live": 25, "interactive": 26, "podcast": 27, "vertical_rows": (28, 29)},
    3: {"date_row": 31, "full_eps": 32, "exclusive": 33, "lifts": 34, "compilations": 35, "live": 36, "interactive": 37, "podcast": 38, "vertical_rows": (39, 40)},
    4: {"date_row": 42, "full_eps": 43, "exclusive": 44, "lifts": 45, "compilations": 46, "live": 47, "interactive": 48, "podcast": 49, "vertical_rows": (50, 51)},
    5: {"date_row": 53, "full_eps": 54, "exclusive": 55, "lifts": 56, "compilations": 57, "live": 58, "interactive": 59, "podcast": 60, "vertical_rows": (61, 62)},
}

DJR_OVERVIEW_FILL_COLORS = {
    "full_eps": "FFFF00", "exclusive": "B4C6E7", "lifts": "A9D08E",
    "compilations": "F4B084", "live": "00B0F0", "podcast": "CC99FF", "vertical": "FF99CC",
    # Nota: Interactive no está aquí, así mantendrá el color original que le hayas puesto al template
}

PROPERTY_ABBREVIATIONS = {
    "Marvel Iron Man and His Awesome Friends": "Marvel Iron Man AHAF",
    "Mickey Mouse Clubhouse": "MMCH+",
    "Mickey Mouse and the Roadster Racers": "MMRR",
    "Mickey Mouse and the Roadster Racers'": "MMRR",
}

def clean_property(value):
    if isinstance(value, list): value = " / ".join(normalize_str(x) for x in value if normalize_str(x))
    value = normalize_str(value).replace("[", "").replace("]", "").replace("'", "").replace('"', "")
    return PROPERTY_ABBREVIATIONS.get(value, value)

def get_episode_category(record):
    category = normalize_str(record.get(AT_FIELDS.get("episode_category", "Episode Category")))
    if not category: return "CURRENT"
    category_upper = category.upper()
    if "LIBRARY" in category_upper: return "LIBRARY"
    if "CURRENT" in category_upper: return "CURRENT"
    return category_upper

def get_djr_overview_region_text(record):
    region_text = normalize_overview_region(record.get(AT_FIELDS["region"]))
    return region_text.replace("TIER A, B & C", "TIER A, B, C")

def get_episode_number(record):
    title = normalize_str(record.get(AT_FIELDS["youtube_title"]))
    activity = normalize_str(record.get(AT_FIELDS["activity"]))
    source = f"{activity} {title}"
    match = re.search(r"\b(?:ep|episode|episodio)\.?\s*(\d+)", source, re.I)
    if match: return match.group(1)
    return ""

def build_djr_overview_text(record, target_country=None):
    content_type = normalize_str(record.get(AT_FIELDS["content_type"]))
    prop = clean_property(record.get(AT_FIELDS["property"]))
    title = normalize_str(record.get(AT_FIELDS["youtube_title"]))
    activity = normalize_str(record.get(AT_FIELDS["activity"]))
    
    # Si es calendario local, no imprimirá el bloque de TIER
    region = "" if target_country else get_djr_overview_region_text(record)
    
    fields = record.get("fields", record)
    policy_val = fields.get("Rights Management Policy")
    policy_str = str(policy_val).strip() if policy_val else ""
    
    gf_suffix = ""
    if "Block AUS/NZ" in policy_str: gf_suffix = " [BLOCK AUS/NZ]"
    elif "Monetize Locally" in policy_str: gf_suffix = " GF"
    elif "Block Everywhere" in policy_str: gf_suffix = " [🚫 NO MONETIZE]"

    ep = get_episode_number(record)

    def _join_parts(*args):
        return "\n".join(str(a).strip() for a in args if str(a).strip())

    if content_type == "Full Episode":
        ep_text = f" ep. {ep}" if ep else ""
        line1 = f"{prop}{ep_text} {title}".strip()
        line2 = get_episode_category(record)
        return _join_parts(line1 + gf_suffix, line2, region)

    if content_type in {"Exclusive", "Editorial Short"}:
        ep_text = f" ep. {ep}" if ep else ""
        return _join_parts(f"{prop} Short{ep_text}", region + gf_suffix)

    if content_type in {"Lift", "Musical Lift", "Magical Moment"}:
        label = "Magical Moment" if content_type == "Magical Moment" else content_type
        return f"{label} {prop}{gf_suffix}".strip()

    if content_type == "Compilation":
        prefix = "MEGA Compilation" if "mega" in activity.lower() or "mega" in title.lower() else "Compilation"
        return _join_parts(f"{prefix} {prop}", region + gf_suffix)

    if content_type == "Live Stream":
        if "full" in title.lower() or "episode" in title.lower(): return _join_parts(f"{prop} Compilation Full Eps", region + gf_suffix)
        return _join_parts(f"MEGA {prop}", region + gf_suffix)

    if content_type == "Podcast":
        return _join_parts(f"{prop} Podcast", region + gf_suffix)

    if content_type == "Interactive":
        return _join_parts(f"{prop} Interactive", region + gf_suffix)

    if content_type == "Vertical Short":
        return prop

    return _join_parts(prop, region + gf_suffix)

def get_djr_overview_bucket(content_type):
    if content_type == "Full Episode": return "full_eps"
    if content_type in {"Exclusive", "Editorial Short"}: return "exclusive"
    if content_type in {"Lift", "Musical Lift", "Magical Moment"}: return "lifts"
    if content_type == "Compilation": return "compilations"
    if content_type == "Live Stream": return "live"
    if content_type == "Podcast": return "podcast"
    if content_type == "Interactive": return "interactive"
    if content_type == "Vertical Short": return "vertical"
    return None

def append_to_cell(cell, text):
    existing = normalize_str(cell.value)
    if existing: cell.value = f"{existing}\n\n— — —\n\n{text}"
    else: cell.value = text

def write_djr_overview_regular_cell(ws, row, col, text, bucket):
    cell = get_writable_cell(ws, row, col)
    append_to_cell(cell, text)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    cell.font = Font(bold=True, color="FF000000")
    fill_color = DJR_OVERVIEW_FILL_COLORS.get(bucket)
    if fill_color: cell.fill = PatternFill(fill_type="solid", fgColor=fill_color)

def write_djr_overview_vertical(ws, layout, col, record, target_country=None):
    top_row, bottom_row = layout["vertical_rows"]
    top_cell = get_writable_cell(ws, top_row, col)
    bottom_cell = get_writable_cell(ws, bottom_row, col)
    top_text = build_djr_overview_text(record, target_country=target_country)
    bottom_text = get_djr_overview_region_text(record)
    append_to_cell(top_cell, top_text)

    fill = PatternFill(fill_type="solid", fgColor=DJR_OVERVIEW_FILL_COLORS["vertical"])
    top_cell.fill = fill
    top_cell.font = Font(bold=True, color="FF000000")
    top_cell.alignment = center_wrap()

    if target_country:
        if isinstance(bottom_cell, Cell):
            bottom_cell.value = None
        bottom_cell.fill = fill 
        try: ws.merge_cells(start_row=top_row, start_column=col, end_row=bottom_row, end_column=col)
        except Exception: pass
    else:
        bottom_cell.value = bottom_text
        bottom_cell.fill = fill
        bottom_cell.alignment = center_wrap()
        if bottom_text == "SPAIN ONLY": bottom_cell.font = Font(bold=True, color="FFFFFF00")
        else: bottom_cell.font = Font(bold=True, color="FFFF0000")

def write_djr_overview_header(ws, records):
    month_name = month_name_from_records(records)
    title_cell = get_writable_cell(ws, 5, 2)
    existing = normalize_str(title_cell.value)
    if existing: title_cell.value = existing.replace("MONTH", month_name)
    else: title_cell.value = f"DJR EMEA YT {month_name} 2026 CALENDAR"
    title_cell.alignment = center_wrap()

    first_monday = first_visible_monday_of_target_month(records)
    if first_monday:
        date_cell = get_writable_cell(ws, 9, 4)
        date_cell.value = first_monday
        date_cell.number_format = "dd-mmm"
        date_cell.alignment = center_wrap()
    return first_monday

def auto_adjust_djr_overview_heights(ws):
    for layout in DJR_OVERVIEW_WEEK_LAYOUT.values():
        rows = [
            layout["full_eps"], layout["exclusive"], layout["lifts"],
            layout["compilations"], layout["live"], layout["interactive"], layout["podcast"], *layout["vertical_rows"],
        ]
        for row in rows:
            ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 15, 42)

def is_country_match(target_country: str, included_list: set) -> bool: 
    if not target_country: return True
    tc = target_country.lower().split("(")[0].strip()
    
    aliases = {
        "GERMANY": ["GSA", "GERMAN"],
        "TÜRKIYE": ["TURKEY", "TUKEY"],
        "SPAIN": ["ESPAÑA"]
    }
    match_list = [tc] + aliases.get(tc.upper(), [])
    
    for inc in included_list:
        ic = str(inc).lower().split("(")[0].strip()
        for t in match_list:
            if t == ic or t in ic or ic in t:
                return True
    return False

def populate_monthly_overview_djr(
    wb, records, enable_mediaplanning=True,
    channel_id_to_name_map=None, target_country=None 
):
    overview_sheet_name = None
    for sn in wb.sheetnames:
        if "overview" in sn.lower() or "monthly" in sn.lower():
            overview_sheet_name = sn
            break
            
    if not overview_sheet_name: return
    ws = wb[overview_sheet_name]

    if target_country:
        thin_side = Side(border_style="thin", color="000000")
        standard_border = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)
        
        for layout in DJR_OVERVIEW_WEEK_LAYOUT.values():
            for col in range(4, 11): # Lunes a Domingo
                # 1. Descombinar SOLO los bloques verticales
                try: ws.unmerge_cells(start_row=layout["vertical_rows"][0], start_column=col, end_row=layout["vertical_rows"][1], end_column=col)
                except: pass

                # 2. Vaciado Quirúrgico
                rows_to_clear = [
                    layout["full_eps"], layout["exclusive"], layout["lifts"],
                    layout["compilations"], layout["live"], 
                    layout["interactive"], layout["podcast"],
                    layout["vertical_rows"][0], layout["vertical_rows"][1]
                ]
                for r in rows_to_clear:
                    try:
                        cell = ws.cell(row=r, column=col)
                        cell.value = None
                        cell.fill = PatternFill(fill_type=None)
                        cell.border = standard_border 
                    except: pass

    first_monday = write_djr_overview_header(ws, records) 
    if not first_monday: return

    for rec in records:
        fields = rec.get("fields", rec)
        if asset_is_event(fields): continue
            
        if target_country:
            from exporters.full_calendar_djr import collect_included_countries_djr
            included = collect_included_countries_djr(record=rec, channel_id_to_name_map=channel_id_to_name_map)
            if not is_country_match(target_country, included): continue

        owner = hub_role(fields.get(AT_FIELDS["owner_hub"]))
        if owner == "Mediaplanning" and not enable_mediaplanning: continue

        live_date = parse_date(fields.get(AT_FIELDS["live_date"]))
        if not live_date: continue

        delta_days = (live_date - first_monday).days
        if delta_days < 0: continue

        week_index = (delta_days // 7) + 1
        if week_index not in DJR_OVERVIEW_WEEK_LAYOUT: continue

        day_name = DAY_ORDER[live_date.weekday()]
        col = DJR_OVERVIEW_DAY_COLS[day_name]
        content_type = normalize_str(fields.get(AT_FIELDS["content_type"]))
        bucket = get_djr_overview_bucket(content_type)
        if not bucket: continue

        layout = DJR_OVERVIEW_WEEK_LAYOUT[week_index]

        if bucket == "vertical":
            write_djr_overview_vertical(ws, layout, col, fields, target_country=target_country)
            continue

        row = layout[bucket]

        text = build_djr_overview_text(fields, target_country=target_country)
        write_djr_overview_regular_cell(ws, row, col, text, bucket)

    auto_adjust_djr_overview_heights(ws)
    logging.info(f"Finished populating '{overview_sheet_name}'.")