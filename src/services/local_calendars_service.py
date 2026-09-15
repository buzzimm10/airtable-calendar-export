import os
import logging
import zipfile
import re
from typing import List, Optional, Set
from openpyxl import load_workbook
from datetime import datetime
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

from config.layouts import TARGET_SHEET as dc_sheet_name
from config.countries import SHEET_COUNTRIES as dc_master_countries
from config.layouts import DAY_ORDER
from exporters.full_calendar_djr import DJR_TARGET_SHEET as djr_sheet_name
from config.countries_djr import DJR_SHEET_COUNTRIES as djr_master_countries
from config.layouts_djr import DJR_WEEK_LAYOUT as djr_layout
from config.layouts_djr import DJR_DAY_LAYOUT as djr_day
from core.excel_utils import safe_set_value

def get_clean_country_name(name: str) -> str:
    return re.sub(r'\s*\(.*?\)', '', name).strip()

def generate_local_calendars(master_path: str, output_dir: str, channel_type: str, 
                             selected_countries: Optional[List[str]] = None,
                             month_str: str = "2026-05", 
                             dc_blocks: Optional[dict] = None,
                             records: Optional[list] = None, 
                             channel_id_to_name_map: Optional[dict] = None):
    
    if not os.path.exists(master_path) or not selected_countries: return None

    os.makedirs(output_dir, exist_ok=True)
    for f in os.listdir(output_dir):
        if f.endswith(".xlsx") or f.endswith(".zip"): os.remove(os.path.join(output_dir, f))

    is_djr = channel_type.upper() == "DJR"
    sheet_name = djr_sheet_name if is_djr else dc_sheet_name
    master_countries = djr_master_countries if is_djr else dc_master_countries
    
    try:
        dt = datetime.strptime(month_str, "%Y-%m")
        month_num = dt.strftime("%m")
        month_name_year = dt.strftime("%B %Y")
    except:
        month_num = "00"; month_name_year = month_str

    brand_label = "DJR" if is_djr else "DC"
    generated_files = []

    fill_grey = PatternFill(fill_type="solid", fgColor="BFBFBF")
    fill_white = PatternFill(fill_type="solid", fgColor="FFFFFF")
    font_bold = Font(bold=True, color="000000")

    thin_side = Side(border_style="thin", color="000000")
    thick_side = Side(border_style="medium", color="000000") 
    border_normal = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)
    border_thick_bottom = Border(top=thin_side, bottom=thick_side, left=thin_side, right=thin_side)

    def style_recycled_row(sheet, row_idx, label, fill_color, is_last=False):
        if not isinstance(row_idx, int): return
        applied_border = border_thick_bottom if is_last else border_normal
        cell_a = sheet.cell(row=row_idx, column=1)
        if isinstance(cell_a, Cell): cell_a.value = None
        cell_c = sheet.cell(row=row_idx, column=3)
        if isinstance(cell_c, Cell): cell_c.value = None
        safe_set_value(sheet, row_idx, 2, label)
        sheet.row_dimensions[row_idx].hidden = False
        sheet.row_dimensions[row_idx].height = 20  
        for c in range(2, 32): 
            cell = sheet.cell(row=row_idx, column=c)
            if c >= 3 and isinstance(cell, Cell): cell.value = None
            cell.fill = fill_color
            cell.border = applied_border
        cell2 = sheet.cell(row=row_idx, column=2)
        cell2.font = font_bold
        cell2.alignment = Alignment(horizontal="right", vertical="center")

    def clear_cell(sheet, row_idx, col_idx):
        if isinstance(row_idx, int) and isinstance(col_idx, int):
            cell_obj = sheet.cell(row=row_idx, column=col_idx)
            if isinstance(cell_obj, Cell): cell_obj.value = None

    def get_val(sheet, r, c):
        if not isinstance(r, int) or not isinstance(c, int): return ""
        v = sheet.cell(row=r, column=c).value
        return str(v).strip().lower() if v is not None else ""

    def format_israel_tx_tc(sheet, row_idx, is_tx=False):
        if not isinstance(row_idx, int): return
        sheet.row_dimensions[row_idx].hidden = False
        sheet.row_dimensions[row_idx].height = 60 if is_tx else 35 
        for c in range(4, 32):
            cell = sheet.cell(row=row_idx, column=c)
            if isinstance(cell, Cell): cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for country in selected_countries:
        if country.lower() in ["spain", "españa"]: continue
        country_wb = load_workbook(master_path, data_only=False)
        clean_country = get_clean_country_name(country)
        
        for sn in list(country_wb.sheetnames):
            sn_lower = sn.lower()
            if "compilation" in sn_lower or "sheet 1" in sn_lower or "sheet1" in sn_lower or "id's" in sn_lower or "metadata" in sn_lower:
                del country_wb[sn]
            
        if sheet_name not in country_wb.sheetnames:
            country_wb.close()
            continue
            
        ws = country_wb[sheet_name]
        is_israel = clean_country.lower() == "israel"
        max_content_row: int = 0
        current_max_row: int = ws.max_row if isinstance(ws.max_row, int) else 500
        safe_rows: Set[int] = set()

        ws.column_dimensions['AG'].hidden = True
        ws.column_dimensions['AF'].hidden = True

        if records:
            if is_djr:
                from exporters.monthly_overview_djr import populate_monthly_overview_djr
                populate_monthly_overview_djr(country_wb, records=records, enable_mediaplanning=True, channel_id_to_name_map=channel_id_to_name_map, target_country=clean_country)
            else:
                from exporters.monthly_overview import populate_monthly_overview
                populate_monthly_overview(country_wb, records=records, enable_mediaplanning=True, channel_id_to_name_map=channel_id_to_name_map, target_country=clean_country)

        # ---------------- LÓGICA DISNEY JUNIOR ----------------
        if is_djr:
            country_idx = next((i for i, mc in enumerate(master_countries) if mc.lower() == country.lower()), -1)
            if country_idx == -1: country_wb.close(); continue
                
            for week_num, layout in djr_layout.items():
                rollout_start = layout["country_rows"][0]
                my_rollout_row = rollout_start + country_idx
                
                if my_rollout_row > 0:
                    safe_rows.add(my_rollout_row)
                    cell_b = ws.cell(row=my_rollout_row, column=2)
                    cell_c = ws.cell(row=my_rollout_row, column=3)
                    country_name_val = cell_c.value if isinstance(cell_c, Cell) else None
                    if isinstance(cell_b, Cell): cell_b.value = country_name_val; cell_b.alignment = Alignment(horizontal='center', vertical='center')
                    if isinstance(cell_c, Cell): cell_c.value = None
                    try: ws.merge_cells(start_row=my_rollout_row, start_column=2, end_row=my_rollout_row, end_column=3)
                    except: pass

                djr_tc = layout.get("tc_row")
                if isinstance(djr_tc, int) and djr_tc > max_content_row: max_content_row = djr_tc
                
                candidate_rows = [r for r in range(rollout_start, rollout_start + len(master_countries)) if r > my_rollout_row]
                if len(candidate_rows) < 3:
                    status_start = layout.get("status_rows", [0])[0]
                    candidate_rows.extend([r for r in range(status_start, status_start + 30) if r > my_rollout_row])
                
                candidate_rows = sorted(list(set(candidate_rows)))
                if len(candidate_rows) >= 3:
                    used_row = candidate_rows[0]
                    link_row = candidate_rows[1]
                    blank_row = candidate_rows[2]
                    
                    style_recycled_row(ws, used_row, "USED OR NOT", fill_grey, is_last=False)
                    style_recycled_row(ws, link_row, "LINK YOUTUBE", fill_white, is_last=True)
                    ws.row_dimensions[blank_row].hidden = False
                    ws.row_dimensions[blank_row].height = 15
                    for c in range(1, 32):
                        cell = ws.cell(row=blank_row, column=c)
                        if isinstance(cell, Cell): cell.value = None
                        cell.fill = PatternFill(fill_type=None)
                        cell.border = Border()
                    safe_rows.update([used_row, link_row, blank_row])

                for day in DAY_ORDER:
                    for role, col in djr_day[day]:
                        cell = ws.cell(row=my_rollout_row, column=col)
                        cell_val = str(cell.value).strip().upper() if hasattr(cell, 'value') else ""
                        if "YES" not in cell_val:
                            if isinstance(cell, Cell): cell.value = None
                            rows_to_clear = [layout["property_row"], layout["content_type_row"], layout["title_row"], layout["synopsis_row"], layout["strategy_row"]]
                            if is_israel: rows_to_clear.extend([layout.get("tx_row"), layout.get("tc_row")])
                            for r in rows_to_clear: clear_cell(ws, r, col)

                if is_israel:
                    israel_tx = layout.get("tx_row")
                    if isinstance(israel_tx, int): format_israel_tx_tc(ws, israel_tx, is_tx=True); safe_rows.add(israel_tx)
                    israel_tc = layout.get("tc_row")
                    if isinstance(israel_tc, int): format_israel_tx_tc(ws, israel_tc, is_tx=False); safe_rows.add(israel_tc)

        # ---------------- LÓGICA DISNEY CHANNEL ----------------
        else:
            if not dc_blocks: country_wb.close(); continue
            for block in dc_blocks.values():
                target_key = next((k for k in block.rollout_country_rows.keys() if k.lower() == country.lower()), None)
                country_row_rollout = block.rollout_country_rows.get(target_key) if target_key else None
                
                if country_row_rollout:
                    safe_rows.add(country_row_rollout)
                    cell_b = ws.cell(row=country_row_rollout, column=2)
                    cell_c = ws.cell(row=country_row_rollout, column=3)
                    country_name_val = cell_c.value if isinstance(cell_c, Cell) else None
                    if isinstance(cell_b, Cell): cell_b.value = country_name_val; cell_b.alignment = Alignment(horizontal='center', vertical='center')
                    if isinstance(cell_c, Cell): cell_c.value = None
                    try: ws.merge_cells(start_row=country_row_rollout, start_column=2, end_row=country_row_rollout, end_column=3)
                    except: pass

                    candidate_rows = [r for r in block.rollout_country_rows.values() if r > country_row_rollout]
                    if len(candidate_rows) < 3: candidate_rows.extend([r for r in block.status_country_rows.values() if r > country_row_rollout])
                    
                    candidate_rows = sorted(list(set(candidate_rows)))
                    if len(candidate_rows) >= 3:
                        used_row = candidate_rows[0]
                        link_row = candidate_rows[1]
                        blank_row = candidate_rows[2]
                        style_recycled_row(ws, used_row, "USED OR NOT", fill_grey, is_last=False)
                        style_recycled_row(ws, link_row, "LINK YOUTUBE", fill_white, is_last=True)
                        ws.row_dimensions[blank_row].hidden = False
                        ws.row_dimensions[blank_row].height = 15
                        for c in range(1, 32):
                            cell = ws.cell(row=blank_row, column=c)
                            if isinstance(cell, Cell): cell.value = None
                            cell.fill = PatternFill(fill_type=None)
                            cell.border = Border()
                        safe_rows.update([used_row, link_row, blank_row])

                dc_tc = block.tc_row
                if isinstance(dc_tc, int) and dc_tc > max_content_row: max_content_row = dc_tc
                
                if country_row_rollout:
                    for col in range(4, 32): 
                        cell = ws.cell(row=country_row_rollout, column=col)
                        cell_val = str(cell.value).strip().upper() if hasattr(cell, 'value') else ""
                        if "YES" not in cell_val:
                            if isinstance(cell, Cell): cell.value = None
                            rows_to_clear = [block.content_type_row, block.property_row, block.title_row, block.synopsis_row, block.strategy_row]
                            if is_israel: rows_to_clear.extend([block.tx_row, block.tc_row])
                            rows_to_clear = [r for r in rows_to_clear if r is not None]
                            for r in rows_to_clear: clear_cell(ws, r, col)

                if is_israel:
                    israel_dc_tx = block.tx_row
                    if isinstance(israel_dc_tx, int): format_israel_tx_tc(ws, israel_dc_tx, is_tx=True); safe_rows.add(israel_dc_tx)
                    israel_dc_tc = block.tc_row
                    if isinstance(israel_dc_tc, int): format_israel_tx_tc(ws, israel_dc_tc, is_tx=False); safe_rows.add(israel_dc_tc)

        for r in range(5, current_max_row + 1):
            if get_val(ws, r, 2).startswith("week"):
                prev_r = r - 1
                safe_rows.add(r)
                safe_rows.add(prev_r)
                ws.row_dimensions[prev_r].hidden = False
                ws.row_dimensions[prev_r].height = 15
                for c in range(1, 32):
                    cell = ws.cell(row=prev_r, column=c)
                    if isinstance(cell, Cell): cell.value = None
                    cell.fill = PatternFill(fill_type=None)
                    cell.border = Border() 

        # ---------------- LA NUEVA GUILLOTINA (Radar Inteligente) ----------------
        rows_to_hide_and_clear = set()
        country_names_lower = [c.lower() for c in master_countries]
        
        standard_headers = [
            "deadline", "owner", "property", "content type", 
            "asset", "title", "synopsis", "planning strategy"
        ]
        
        junk_keywords = [
            "watermark", "bug", "copyright", "endcard", "branding", 
            "status conforming", "comments", "total vod", "burned"
        ]

        for r in range(1, current_max_row + 1):
            val_a = get_val(ws, r, 1)
            val_b = get_val(ws, r, 2)
            val_c = get_val(ws, r, 3)
            
            # EL FIX: Solo leemos las 3 primeras columnas para no ocultar "Ladybugs"
            header_col_text = f"{val_a} {val_b} {val_c}".lower()
            
            if any(h in header_col_text for h in standard_headers) or "week" in header_col_text:
                safe_rows.add(r)
                continue

            if r in safe_rows:
                continue
            
            if any(k in header_col_text for k in junk_keywords):
                rows_to_hide_and_clear.add(r)
                continue
                
            if val_b.startswith("tier ") or val_a.startswith("tier "):
                rows_to_hide_and_clear.add(r)
                continue

            if any(c in header_col_text for c in country_names_lower) or "total" in header_col_text:
                rows_to_hide_and_clear.add(r)
                continue

        if max_content_row > 0:
            for r in range(max_content_row + 5, current_max_row + 1): 
                if r not in safe_rows:
                    rows_to_hide_and_clear.add(r)

        # ---------------- EJECUCIÓN (VACIADO Y OCULTACIÓN SEGURA) ----------------
        for r in sorted(list(rows_to_hide_and_clear)):
            ws.row_dimensions[r].hidden = True
            ws.row_dimensions[r].height = 0
            for c in range(1, 32):
                cell = ws.cell(row=r, column=c)
                if isinstance(cell, Cell): cell.value = None
                cell.fill = PatternFill(fill_type=None)
                cell.border = Border()

        # --- REEMPLAZO DE NOMBRES EMEA -> COUNTRY ---
        for sn in country_wb.sheetnames:
            ws_tmp = country_wb[sn]
            for r in range(1, 15):
                for c in range(1, 10):
                    c_obj = ws_tmp.cell(row=r, column=c)
                    if isinstance(c_obj, Cell):
                        v = c_obj.value
                        if isinstance(v, str) and "EMEA" in v: c_obj.value = v.replace("EMEA", clean_country.upper())
        
        for sn in country_wb.sheetnames:
            if "EMEA" in sn:
                new_title = sn.replace("EMEA", clean_country.upper())
                country_wb[sn].title = new_title[:31] 

        filename = f"{month_num}_Social Media Calendar {brand_label} YouTube {clean_country} - {month_name_year}.xlsx"
        output_path = os.path.join(output_dir, filename)
        country_wb.template = False
        country_wb.save(output_path)
        country_wb.close()
        generated_files.append(output_path)

    zip_filename = f"Local_Calendars_{brand_label}_{month_num}.zip"
    zip_path = os.path.join(output_dir, zip_filename)
    with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
        for f in generated_files: zipf.write(f, os.path.basename(f))
            
    return zip_path