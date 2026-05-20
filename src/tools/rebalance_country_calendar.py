import argparse
import json
import os
import shutil
from copy import copy
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter


WEEK_START_ROWS = [6, 17, 28, 39, 50]

ROW_OFFSETS = {
    "week_day": 0,
    "owner": 1,
    "property": 2,
    "content_type": 3,
    "title": 4,
    "synopsis": 5,
    "strategy": 6,
    "used_asset": 7,
    "yt_link": 8,
}

ASSET_ROW_KEYS = [
    "owner",
    "property",
    "content_type",
    "title",
    "synopsis",
    "strategy",
    "used_asset",
    "yt_link",
]

DAY_COLUMNS = {
    "Monday": [4, 5, 6, 7],       # D:G
    "Tuesday": [8, 9, 10, 11],    # H:K
    "Wednesday": [12, 13, 14, 15],# L:O
    "Thursday": [16, 17, 18, 19], # P:S
    "Friday": [20, 21, 22, 23],   # T:W
    "Saturday": [24, 25, 26, 27], # X:AA
    "Sunday": [28, 29, 30, 31],   # AB:AE
}

DAY_ORDER = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]

EXPECTED_COUNTRY_FILES = [
    "6_Social Media Calendar DC YouTube EMEA - June 2026 - Belgique.xlsx",
    "6_Social Media Calendar DC YouTube EMEA - June 2026 - Bulgaria.xlsx",
    "6_Social Media Calendar DC YouTube EMEA - June 2026 - czech R.xlsx",
    "6_Social Media Calendar DC YouTube EMEA - June 2026 - denmark.xlsx",
    "6_Social Media Calendar DC YouTube EMEA - June 2026 - FR.xlsx",
    "6_Social Media Calendar DC YouTube EMEA - June 2026 - Greece.xlsx",
    "6_Social Media Calendar DC YouTube EMEA - June 2026 - Hungary.xlsx",
    "6_Social Media Calendar DC YouTube EMEA - June 2026 - Italy.xlsx",
    "6_Social Media Calendar DC YouTube EMEA - June 2026 - MENA.xlsx",
    "6_Social Media Calendar DC YouTube EMEA - June 2026 - nederlands.xlsx",
    "6_Social Media Calendar DC YouTube EMEA - June 2026 - norway.xlsx",
    "6_Social Media Calendar DC YouTube EMEA - June 2026 - Poland.xlsx",
    "6_Social Media Calendar DC YouTube EMEA - June 2026 - Portugal.xlsx",
    "6_Social Media Calendar DC YouTube EMEA - June 2026 - Romania.xlsx",
    "6_Social Media Calendar DC YouTube EMEA - June 2026 - sweden.xlsx",
]

CONTENT_TYPE_ALIASES = {
    "compilado": "compilation",
    "compilation": "compilation",
    "compilación": "compilation",
    "compilacion": "compilation",

    "episodio completo": "full episode",
    "full episode": "full episode",
    "full episodes": "full episode",

    "sneak peek": "lift",
    "lift": "lift",

    "live stream": "live stream",
    "livestream": "live stream",
}


def normalize_text(value) -> str:
    return str(value or "").strip()


def normalize_content_type(value) -> str:
    raw = normalize_text(value).lower()
    return CONTENT_TYPE_ALIASES.get(raw, raw)


def parse_iso_date(value: str) -> date:
    return datetime.strptime(value, "%Y-%m-%d").date()


def monday_of_week(d: date) -> date:
    return d - timedelta(days=d.weekday())


def get_visible_first_monday(target_date: date) -> date:
    return monday_of_week(date(target_date.year, target_date.month, 1))


def visual_position(d: date, first_monday: date):
    delta_days = (d - first_monday).days
    if delta_days < 0:
        raise ValueError(f"Date {d} is before first visible Monday {first_monday}")

    week_index = (delta_days // 7)
    day_index = delta_days % 7

    if week_index < 0 or week_index >= len(WEEK_START_ROWS):
        raise ValueError(f"Date {d} is outside the 5-week template range")

    return week_index, DAY_ORDER[day_index]


def is_writable_cell(ws, row: int, col: int) -> bool:
    return not isinstance(ws.cell(row=row, column=col), MergedCell)


def get_real_cell(ws, row: int, col: int):
    cell = ws.cell(row=row, column=col)

    if not isinstance(cell, MergedCell):
        return cell

    for merged_range in ws.merged_cells.ranges:
        if (
            merged_range.min_row <= row <= merged_range.max_row
            and merged_range.min_col <= col <= merged_range.max_col
        ):
            return ws.cell(row=merged_range.min_row, column=merged_range.min_col)

    raise RuntimeError(f"Could not resolve merged cell {get_column_letter(col)}{row}")


def copy_cell(src, dst):
    dst.value = src.value

    if src.has_style:
        dst._style = copy(src._style)

    dst.number_format = src.number_format
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.protection = copy(src.protection)


def clear_cell_keep_style(cell):
    cell.value = None


def read_asset_block(ws, week_start_row: int, col: int):
    block = {}

    for key in ASSET_ROW_KEYS:
        row = week_start_row + ROW_OFFSETS[key]
        cell = get_real_cell(ws, row, col)
        block[key] = {
            "value": cell.value,
            "style": copy(cell._style) if cell.has_style else None,
            "number_format": cell.number_format,
            "font": copy(cell.font),
            "fill": copy(cell.fill),
            "border": copy(cell.border),
            "alignment": copy(cell.alignment),
            "protection": copy(cell.protection),
        }

    return block


def write_asset_block(ws, week_start_row: int, col: int, block):
    for key in ASSET_ROW_KEYS:
        row = week_start_row + ROW_OFFSETS[key]
        dst = get_real_cell(ws, row, col)
        data = block[key]

        dst.value = data["value"]

        if data["style"] is not None:
            dst._style = copy(data["style"])

        dst.number_format = data["number_format"]
        dst.font = copy(data["font"])
        dst.fill = copy(data["fill"])
        dst.border = copy(data["border"])
        dst.alignment = copy(data["alignment"])
        dst.protection = copy(data["protection"])


def clear_asset_block(ws, week_start_row: int, col: int):
    for key in ASSET_ROW_KEYS:
        row = week_start_row + ROW_OFFSETS[key]
        cell = get_real_cell(ws, row, col)
        clear_cell_keep_style(cell)


def cell_has_asset(ws, week_start_row: int, col: int) -> bool:
    title = normalize_text(get_real_cell(ws, week_start_row + ROW_OFFSETS["title"], col).value)
    prop = normalize_text(get_real_cell(ws, week_start_row + ROW_OFFSETS["property"], col).value)
    content_type = normalize_text(get_real_cell(ws, week_start_row + ROW_OFFSETS["content_type"], col).value)

    return bool(title or prop or content_type)


def get_content_type_at(ws, week_start_row: int, col: int) -> str:
    value = get_real_cell(ws, week_start_row + ROW_OFFSETS["content_type"], col).value
    return normalize_content_type(value)


def find_asset_col(ws, target_date: date, content_type: str, slot=None):
    first_monday = get_visible_first_monday(target_date)
    week_idx, day_name = visual_position(target_date, first_monday)

    week_start_row = WEEK_START_ROWS[week_idx]
    day_cols = DAY_COLUMNS[day_name]

    desired_type = normalize_content_type(content_type)

    matching_cols = [
        col for col in day_cols
        if get_content_type_at(ws, week_start_row, col) == desired_type
    ]

    if slot is not None:
        slot_index = int(slot) - 1
        if slot_index < 0 or slot_index >= len(day_cols):
            return None, week_start_row, day_name

        col = day_cols[slot_index]
        if get_content_type_at(ws, week_start_row, col) == desired_type:
            return col, week_start_row, day_name

        return None, week_start_row, day_name

    if matching_cols:
        return matching_cols[0], week_start_row, day_name

    return None, week_start_row, day_name


def find_destination_col(ws, target_date: date, content_type: str, preferred_role="Programming"):
    first_monday = get_visible_first_monday(target_date)
    week_idx, day_name = visual_position(target_date, first_monday)

    week_start_row = WEEK_START_ROWS[week_idx]
    day_cols = DAY_COLUMNS[day_name]

    programming_cols = day_cols[:2]
    media_cols = day_cols[2:3]
    uk_cols = day_cols[3:4]

    desired_type = normalize_content_type(content_type)

    # 1. If same content type already exists, overwrite that slot
    for col in programming_cols:
        if get_content_type_at(ws, week_start_row, col) == desired_type:
            return col, week_start_row, day_name

    # 2. Prefer Programming slot 1 if empty
    if not cell_has_asset(ws, week_start_row, programming_cols[0]):
        return programming_cols[0], week_start_row, day_name

    # 3. If slot 1 is occupied, use Programming slot 2 if empty
    if not cell_has_asset(ws, week_start_row, programming_cols[1]):
        return programming_cols[1], week_start_row, day_name

    # 4. If both Programming slots are full, overwrite slot 1
    return programming_cols[0], week_start_row, day_name

def compact_programming_slots(ws, log_lines):
    """
    If Programming slot 1 is empty and slot 2 has content, move slot 2 into slot 1.
    Applies across all weeks and days.
    """
    for week_start_row in WEEK_START_ROWS:
        for day_name, day_cols in DAY_COLUMNS.items():
            slot_1 = day_cols[0]
            slot_2 = day_cols[1]

            slot_1_has_asset = cell_has_asset(ws, week_start_row, slot_1)
            slot_2_has_asset = cell_has_asset(ws, week_start_row, slot_2)

            if not slot_1_has_asset and slot_2_has_asset:
                block = read_asset_block(ws, week_start_row, slot_2)
                write_asset_block(ws, week_start_row, slot_1, block)
                clear_asset_block(ws, week_start_row, slot_2)

                log_lines.append(
                    f"COMPACT {day_name} row {week_start_row}: "
                    f"{get_column_letter(slot_2)} -> {get_column_letter(slot_1)}"
                )

def apply_move(ws, move, protected_destinations, log_lines):
    content_type = move["content_type"]
    from_date = parse_iso_date(move["from"])
    to_date = parse_iso_date(move["to"])
    slot = move.get("slot")

    src_col, src_week_row, src_day = find_asset_col(
        ws=ws,
        target_date=from_date,
        content_type=content_type,
        slot=slot,
    )

    if src_col is None:
        log_lines.append(f"MISS MOVE source not found: {content_type} {from_date} slot={slot}")
        return

    dst_col, dst_week_row, dst_day = find_destination_col(
        ws=ws,
        target_date=to_date,
        content_type=content_type,
    )

    block = read_asset_block(ws, src_week_row, src_col)

    write_asset_block(ws, dst_week_row, dst_col, block)
    clear_asset_block(ws, src_week_row, src_col)

    protected_destinations.add((dst_week_row, dst_col))

    log_lines.append(
        f"MOVE {content_type}: {from_date} {src_day} {get_column_letter(src_col)} "
        f"-> {to_date} {dst_day} {get_column_letter(dst_col)}"
    )


def apply_delete(ws, delete, protected_destinations, log_lines):
    content_type = delete["content_type"]
    target_date = parse_iso_date(delete["date"])
    slot = delete.get("slot")

    col, week_row, day_name = find_asset_col(
        ws=ws,
        target_date=target_date,
        content_type=content_type,
        slot=slot,
    )

    if col is None:
        log_lines.append(f"MISS DELETE target not found: {content_type} {target_date} slot={slot}")
        return

    if (week_row, col) in protected_destinations and not delete.get("force", False):
        log_lines.append(
            f"SKIP DELETE protected moved asset: {content_type} {target_date} "
            f"{day_name} {get_column_letter(col)}. Use force=true if intentional."
        )
        return

    clear_asset_block(ws, week_row, col)

    log_lines.append(
        f"DELETE {content_type}: {target_date} {day_name} {get_column_letter(col)}"
    )

def apply_moves_safe(ws, moves, protected_destinations, log_lines):
    planned_moves = []

    # Phase 1: read all original sources before changing anything
    for move in moves:
        content_type = move["content_type"]
        from_date = parse_iso_date(move["from"])
        to_date = parse_iso_date(move["to"])
        slot = move.get("slot")

        src_col, src_week_row, src_day = find_asset_col(
            ws=ws,
            target_date=from_date,
            content_type=content_type,
            slot=slot,
        )

        if src_col is None:
            log_lines.append(f"MISS MOVE source not found: {content_type} {from_date} slot={slot}")
            continue

        dst_col, dst_week_row, dst_day = find_destination_col(
            ws=ws,
            target_date=to_date,
            content_type=content_type,
        )

        block = read_asset_block(ws, src_week_row, src_col)

        planned_moves.append({
            "content_type": content_type,
            "from_date": from_date,
            "to_date": to_date,
            "src_col": src_col,
            "src_week_row": src_week_row,
            "src_day": src_day,
            "dst_col": dst_col,
            "dst_week_row": dst_week_row,
            "dst_day": dst_day,
            "block": block,
        })

    # Phase 2: clear original sources
    for item in planned_moves:
        clear_asset_block(ws, item["src_week_row"], item["src_col"])

    # Phase 3: write all destinations from original snapshots
    for item in planned_moves:
        write_asset_block(ws, item["dst_week_row"], item["dst_col"], item["block"])
        protected_destinations.add((item["dst_week_row"], item["dst_col"]))

        log_lines.append(
            f"MOVE {item['content_type']}: {item['from_date']} {item['src_day']} {get_column_letter(item['src_col'])} "
            f"-> {item['to_date']} {item['dst_day']} {get_column_letter(item['dst_col'])}"
        )

def process_workbook(input_path: Path, output_path: Path, config: dict, sheet_name=None):
    wb = load_workbook(input_path)

    if sheet_name:
        if sheet_name not in wb.sheetnames:
            raise RuntimeError(f"Sheet '{sheet_name}' not found in {input_path.name}")
        ws = wb[sheet_name]
    else:
        ws = wb.active

    log_lines = [f"FILE {input_path.name}"]
    protected_destinations = set()

    apply_moves_safe(
        ws=ws,
        moves=config.get("moves", []),
        protected_destinations=protected_destinations,
        log_lines=log_lines,
)

    for delete in config.get("deletes", []):
        apply_delete(ws, delete, protected_destinations, log_lines)

    compact_programming_slots(ws, log_lines)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    return log_lines


def load_config(config_path: Path):
    with open(config_path, "r", encoding="utf-8") as f:
        return json.load(f)


def process_folder(input_folder: Path, output_folder: Path, config: dict, sheet_name=None):
    all_logs = []

    for filename in EXPECTED_COUNTRY_FILES:
        input_path = input_folder / filename

        if not input_path.exists():
            all_logs.append(f"MISS FILE {filename}")
            continue

        output_path = output_folder / filename

        try:
            logs = process_workbook(
                input_path=input_path,
                output_path=output_path,
                config=config,
                sheet_name=sheet_name,
            )
            all_logs.extend(logs)
            all_logs.append(f"SAVED {output_path}")
            all_logs.append("")
        except Exception as e:
            all_logs.append(f"ERROR {filename}: {e}")
            all_logs.append("")

    return all_logs


def main():
    parser = argparse.ArgumentParser(description="Rebalance country-level DC calendars after planning changes.")

    parser.add_argument("--input", help="Single input .xlsx file")
    parser.add_argument("--output", help="Single output .xlsx file")
    parser.add_argument("--input-folder", help="Folder containing country .xlsx files")
    parser.add_argument("--output-folder", help="Folder for processed output files")
    parser.add_argument("--config", required=True, help="JSON config with moves/deletes")
    parser.add_argument("--sheet", help="Optional sheet name. Defaults to active sheet.")

    args = parser.parse_args()

    config = load_config(Path(args.config))

    if args.input and args.output:
        logs = process_workbook(
            input_path=Path(args.input),
            output_path=Path(args.output),
            config=config,
            sheet_name=args.sheet,
        )
    elif args.input_folder and args.output_folder:
        logs = process_folder(
            input_folder=Path(args.input_folder),
            output_folder=Path(args.output_folder),
            config=config,
            sheet_name=args.sheet,
        )
    else:
        raise RuntimeError("Use either --input/--output or --input-folder/--output-folder")

    log_path = Path(args.output_folder or Path(args.output).parent) / "rebalance_log.txt"
    log_path.parent.mkdir(parents=True, exist_ok=True)

    with open(log_path, "w", encoding="utf-8") as f:
        f.write("\n".join(logs))

    print("\n".join(logs))
    print(f"\nLog saved to: {log_path}")


if __name__ == "__main__":
    main()